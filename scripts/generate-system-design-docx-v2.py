#!/usr/bin/env python3
"""Generate the V2.10 enterprise system design specification.

Every table is cloned from the supplied template. The generator never creates
a table with python-docx's add_table API.
"""

from __future__ import annotations

from copy import deepcopy
from datetime import date
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.shared import Cm, Pt, RGBColor

from assets.system_design_v2_content import (
    FUNCTION_DOMAINS,
    MODULES,
    NONFUNCTIONAL,
    TECHNOLOGY,
    UI_SECTIONS,
    UML_LIST,
    UNCONFIRMED,
)


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "scripts/assets/system-design-reference.docx"
OUTPUT = ROOT / "docs/校园漏洞响应与安全学习平台项目系统设计说明书（重编版）.docx"
SCHEMA = ROOT / "database/prisma/schema.prisma"
BACKEND = ROOT / "backend/src"
SYSTEM_DESIGN_DIAGRAM_DIR = ROOT / "docs/assets/diagrams/system-design"
PROJECT = "校园漏洞响应与安全学习平台"
VERSION = "V2.10"
DOCUMENT_DATE = date(2026, 8, 26)

OVERVIEW_DIAGRAMS = {
    "2-1": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-2-1-system-use-cases.png", 14.2),
    "3-1": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-3-1-deployment.png", 15.3),
    "3-2": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-3-2-component-architecture.png", 15.3),
    "3-3": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-3-3-backend-package-dependencies.png", 15.3),
    "3-4": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-3-4-core-domain-model.png", 15.3),
    "3-5": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-3-5-security-architecture.png", 14.2),
    "5-1": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-5-1-core-er.png", 15.3),
}

MODULE_DIAGRAMS = {
    "4-1": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-1-auth-token-management.png", 15.3, "身份认证与令牌管理时序图"),
    "4-2": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-2-security-admission.png", 15.3, "账户资料与安全准入活动图"),
    "4-3": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-3-vulnerability-draft-submit.png", 15.3, "漏洞草稿与 Markdown 报告提交时序图"),
    "4-4": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-4-attachment-access.png", 15.3, "附件上传与授权访问时序图"),
    "4-5": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-5-vulnerability-lifecycle.png", 15.3, "漏洞审核、补充、复测与归档状态图"),
    "4-6": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-6-certificate-verification.png", 15.3, "证书签发、撤销与公开核验时序图"),
    "4-7": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-7-points-redemption.png", 15.3, "积分商城并发兑换活动图"),
    "4-8": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-8-learning-announcement.png", 15.3, "学习资源与公告内容发布活动图"),
    "4-9": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-9-notification-delivery.png", 15.3, "通知、邮件与菜单未读发布时序图"),
    "4-10": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-10-chat-realtime.png", 15.3, "受控聊天发送与实时发布时序图"),
    "4-11": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-11-ai-runtime-routing.png", 15.3, "AI 模型策略、BYOK 与运行时序图"),
    "4-12": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-12-ai-report-advice.png", 15.3, "漏洞报告解析与按需 AI 建议时序图"),
    "4-13": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-13-knowledge-rag.png", 15.3, "AI 知识沉淀与检索增强活动图"),
    "4-14": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-14-skill-lifecycle.png", 15.3, "声明式 Skill 生命周期与固定关系图"),
    "4-15": (SYSTEM_DESIGN_DIAGRAM_DIR / "figure-4-15-governance-audit.png", 15.3, "用户、系统设置、仪表盘与审计组件图"),
}

PROTOTYPE = {
    "cover": 0,
    "revision": 1,
    "review": 2,
    "distribution": 3,
    "five": 4,
    "glossary": 5,
    "three": 6,
    "nonfunctional": 7,
    "environment": 8,
    "physical": 9,
}

CORE_MODELS = [
    "User", "PlatformNotification", "Role",
    "SecurityExamAttempt", "ResponsibilityAgreement",
    "VulnerabilityReport", "VulnerabilitySubmissionDraft", "VulnerabilitySupplement",
    "VulnerabilityAttachment", "VulnerabilityEvent", "ReviewAssignment", "ReviewComment", "RetestRecord",
    "Certificate", "PointLedgerEntry", "MallProduct", "MallRedemption", "AiAgentRun", "AiSkill",
    "AiKnowledgeEntry", "ChatConversation", "ChatMessage", "AuditLog", "SystemSetting",
]


def style_named(doc, name):
    for style in doc.styles:
        if style.name == name:
            return style
    raise KeyError(name)


def set_run_font(run, name="宋体", size=12, bold=None, color=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def remove_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc):
    normal = style_named(doc, "Normal")
    normal.font.name = "宋体"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size, before, after in ((1, 16, 18, 10), (2, 14, 12, 6), (3, 12, 8, 4)):
        style = style_named(doc, f"Heading {level}")
        style.font.name = "黑体"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style._element.get_or_add_pPr()
        numbering = p_pr.find(qn("w:numPr"))
        if numbering is not None:
            p_pr.remove(numbering)


def configure_page(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)


def clear_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for container in (section.header, section.footer):
        for paragraph in list(container.paragraphs):
            paragraph._element.getparent().remove(paragraph._element)


def add_field(paragraph, instruction, fallback=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hp.add_run(f"{PROJECT}系统设计说明书"), "宋体", 9)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("第 "), "宋体", 9)
    add_field(fp, "PAGE", "1")
    set_run_font(fp.add_run(" 页"), "宋体", 9)
    page_type = section._sectPr.find(qn("w:pgNumType"))
    if page_type is None:
        page_type = OxmlElement("w:pgNumType")
        section._sectPr.append(page_type)
    page_type.set(qn("w:start"), "1")


def heading(doc, text, level):
    p = doc.add_paragraph()
    p.style = style_named(doc, f"Heading {level}")
    set_run_font(p.add_run(text), "黑体", {1: 16, 2: 14, 3: 12}[level], True)
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24 if indent else 0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run_font(p.add_run(text), "宋体", 12)
    return p


def bullets(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        set_run_font(p.add_run(f"（{index}）{item}"), "宋体", 12)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _replace_paragraph_text(paragraph, text):
    p = paragraph._p
    sample_r_pr = None
    for run in p.findall(qn("w:r")):
        r_pr = run.find(qn("w:rPr"))
        if r_pr is not None:
            sample_r_pr = deepcopy(r_pr)
            break
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    r = OxmlElement("w:r")
    if sample_r_pr is not None:
        r.append(sample_r_pr)
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        if index:
            r.append(OxmlElement("w:br"))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = line
        r.append(t)
    p.append(r)
    p_pr = p.get_or_add_pPr()
    numbering = p_pr.find(qn("w:numPr"))
    if numbering is not None:
        p_pr.remove(numbering)
    indentation = p_pr.find(qn("w:ind"))
    if indentation is None:
        indentation = OxmlElement("w:ind")
        p_pr.append(indentation)
    for attr in ("left", "right", "start", "end", "firstLine", "firstLineChars"):
        indentation.set(qn(f"w:{attr}"), "0")
    for attr in ("hanging", "hangingChars"):
        indentation.attrib.pop(qn(f"w:{attr}"), None)


def _replace_cell_text(cell, text):
    paragraphs = cell.paragraphs
    if not paragraphs:
        p = OxmlElement("w:p")
        cell._tc.append(p)
        paragraphs = cell.paragraphs
    _replace_paragraph_text(paragraphs[0], text)
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)


def _prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _allow_automatic_row_height(row):
    tr_pr = row._tr.get_or_add_trPr()
    for height in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(height)


def _set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _format_table(table, has_header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "autofit")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{name}"))
        if border is None:
            border = OxmlElement(f"w:{name}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "000000")
        border.set(qn("w:space"), "0")

    for row_index, row in enumerate(table.rows):
        _allow_automatic_row_height(row)
        _prevent_row_split(row)
        for cell_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.left_indent = Pt(0)
                paragraph.paragraph_format.right_indent = Pt(0)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_index == 0 and has_header) or cell_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, "宋体", 10.5, bold=(row_index == 0 and has_header))


def _set_table_column_widths(table, ratios, total_width_cm=15.8):
    """Set stable column proportions while retaining content-driven row heights."""
    if len(ratios) != len(table.columns) or sum(ratios) <= 0:
        raise ValueError("column width ratios must match the table column count")
    normalized = [ratio / sum(ratios) for ratio in ratios]
    table.autofit = False
    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table._tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for column, ratio in zip(table.columns, normalized):
        width = Cm(total_width_cm * ratio)
        column.width = width
        for cell in column.cells:
            cell.width = width


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def append_cloned_table(doc, reference, prototype_name, rows, has_header=True):
    """Append a populated deepcopy of an original template table."""
    source = reference.tables[PROTOTYPE[prototype_name]]
    table_xml = deepcopy(source._tbl)
    # Legacy Word tables are positioned as floating objects. Preserve their
    # visual XML, but make them inline so long cloned tables cannot overlap
    # headers, captions, or following sections during repagination.
    table_properties = table_xml.tblPr
    floating = table_properties.find(qn("w:tblpPr"))
    if floating is not None:
        table_properties.remove(floating)
    table = Table(table_xml, doc)
    target_count = len(rows) + (1 if has_header else 0)
    while len(table.rows) < target_count:
        sample_index = 1 if has_header and len(table.rows) > 1 else len(table.rows) - 1
        table._tbl.append(deepcopy(table.rows[sample_index]._tr))
    while len(table.rows) > target_count:
        table._tbl.remove(table.rows[-1]._tr)
    start = 1 if has_header else 0
    for row_index, values in enumerate(rows, start):
        cells = table.rows[row_index].cells
        if len(values) != len(cells):
            raise ValueError(f"prototype {prototype_name} expects {len(cells)} values, got {len(values)}")
        for cell, value in zip(cells, values):
            _replace_cell_text(cell, value)
        _prevent_row_split(table.rows[row_index])
    if has_header:
        _repeat_header(table.rows[0])
    _format_table(table, has_header=has_header)
    body_xml = doc._element.body
    body_xml.insert(len(body_xml) - 1, table_xml)
    return table


def table_caption(doc, number, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    prefix = number if str(number).startswith("附表") else f"表{number}"
    set_run_font(p.add_run(f"{prefix}  {title}"), "宋体", 10.5)


def add_template_table(doc, reference, prototype, number, title, headers, rows, column_widths=None):
    table_caption(doc, number, title)
    table = append_cloned_table(doc, reference, prototype, rows, has_header=True)
    if len(headers) != len(table.rows[0].cells):
        raise ValueError(f"header mismatch for {title}")
    for cell, value in zip(table.rows[0].cells, headers):
        _replace_cell_text(cell, value)
    _format_table(table, has_header=True)
    if column_widths is not None:
        _set_table_column_widths(table, column_widths)
    return table


def add_uml_placeholder(doc, number, title, description):
    rendered = OVERVIEW_DIAGRAMS.get(str(number)) or MODULE_DIAGRAMS.get(str(number))
    if rendered:
        path, width_cm, *caption_override = rendered
        caption_title = caption_override[0] if caption_override else title
        if not path.exists():
            raise FileNotFoundError(f"rendered diagram not found: {path}")
        figure = doc.add_paragraph()
        figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.first_line_indent = Pt(0)
        figure.paragraph_format.keep_together = True
        figure.paragraph_format.keep_with_next = True
        figure.paragraph_format.space_before = Pt(8)
        figure.paragraph_format.space_after = Pt(2)
        figure.add_run().add_picture(str(path), width=Cm(width_cm))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.first_line_indent = Pt(0)
        caption.paragraph_format.space_after = Pt(6)
        set_run_font(caption.add_run(f"图{number}  {caption_title}"), "宋体", 10.5)
        return
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "808080")
        node.set(qn("w:space"), "8")
        borders.append(node)
    p_pr.append(borders)
    set_run_font(p.add_run("【UML工程图待绘制】\n"), "黑体", 11, True)
    set_run_font(p.add_run(description), "宋体", 10.5, False, "444444")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_after = Pt(6)
    set_run_font(caption.add_run(f"图{number}  {title}"), "宋体", 10.5)


def add_screenshot_placeholder(doc, number, title, details):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFFFFF")
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "dashed")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "808080")
        node.set(qn("w:space"), "8")
        borders.append(node)
    p_pr.append(borders)
    set_run_font(p.add_run("【界面截图预留位置】\n"), "黑体", 11, True)
    set_run_font(p.add_run(details), "宋体", 10.5, False, "444444")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_after = Pt(6)
    set_run_font(caption.add_run(f"图{number}  {title}"), "宋体", 10.5)


def parse_schema():
    source = SCHEMA.read_text(encoding="utf-8")
    enums = {}
    for match in re.finditer(r"^enum\s+(\w+)\s*\{(.*?)^\}", source, re.M | re.S):
        values = [line.strip().split()[0] for line in match.group(2).splitlines() if line.strip() and not line.strip().startswith("//")]
        enums[match.group(1)] = values
    models = {}
    for match in re.finditer(r"^model\s+(\w+)\s*\{(.*?)^\}", source, re.M | re.S):
        fields, directives = [], []
        for raw in match.group(2).splitlines():
            line = raw.strip()
            if not line or line.startswith("//"):
                continue
            if line.startswith("@@"):
                directives.append(line)
                continue
            parts = line.split()
            if len(parts) >= 2:
                fields.append((parts[0], parts[1], " ".join(parts[2:])))
        models[match.group(1)] = {"fields": fields, "directives": directives}
    return enums, models


def parse_endpoints():
    endpoints = []
    for path in sorted(BACKEND.rglob("*.controller.ts")):
        source = path.read_text(encoding="utf-8")
        controller = re.search(r"@Controller\((?:'([^']*)'|\"([^\"]*)\"|)\)", source)
        prefix = next((group for group in controller.groups() if group is not None), "") if controller else ""
        lines = source.splitlines()
        decorators = []
        for line in lines:
            stripped = line.strip()
            if stripped.startswith("@"):
                decorators.append(stripped)
                continue
            method_match = re.match(r"(?:async\s+)?(\w+)\s*\(", stripped)
            if not method_match or not decorators:
                if stripped and not stripped.startswith("//") and not stripped.startswith("constructor") and not stripped.startswith("private") and not stripped.startswith("public"):
                    if not stripped.endswith(",") and not stripped.endswith("{"):
                        decorators = []
                continue
            http = None
            route = ""
            permissions = []
            for decorator in decorators:
                hit = re.match(r"@(Get|Post|Put|Patch|Delete)\((?:'([^']*)'|\"([^\"]*)\"|)\)", decorator)
                if hit:
                    http = hit.group(1).upper()
                    route = hit.group(2) or hit.group(3) or ""
                if decorator.startswith("@RequirePermissions"):
                    permissions.extend(re.findall(r"['\"]([^'\"]+)['\"]", decorator))
            if http:
                full = "/api/" + "/".join(part.strip("/") for part in (prefix, route) if part.strip("/"))
                endpoints.append({
                    "method": http,
                    "path": full,
                    "handler": method_match.group(1),
                    "permissions": "、".join(permissions) if permissions else "登录身份或公开接口，以守卫配置为准",
                    "source": str(path.relative_to(ROOT)),
                })
            decorators = []
    return endpoints


FIELD_MEANINGS = {
    "id": "记录主键，使用 cuid 生成", "createdAt": "记录创建时间", "updatedAt": "记录最后更新时间",
    "deletedAt": "软删除时间；为空表示未删除", "userId": "关联用户主键", "reportId": "关联漏洞报告主键",
    "status": "当前生命周期状态", "name": "对象名称", "title": "标题", "description": "说明文本", "summary": "供检索和界面展示使用的内容摘要",
    "content": "正文内容", "markdown": "Markdown 正文", "email": "规范化后的联系邮箱",
    "unifiedCode": "统一认证码，作为校内身份标识", "passwordHash": "密码派生摘要，不保存明文密码",
    "roleId": "关联角色主键", "permissionId": "关联权限主键", "code": "稳定业务编码",
    "version": "内容或并发控制版本号", "expiresAt": "失效时间", "revokedAt": "撤销时间",
    "readAt": "用户读取通知或消息的时间", "ipAddress": "经可信代理链解析的客户端 IP 地址",
    "userAgent": "客户端 User-Agent，用于安全审计", "requestPath": "触发操作的 API 请求路径",
    "metadata": "结构化扩展元数据，禁止保存无必要秘密", "note": "操作备注", "reason": "状态变化或操作原因",
    "objectKey": "S3 兼容对象存储中的私有对象键", "fileName": "用户上传时的原始文件名",
    "mimeType": "经服务端识别的媒体类型", "sizeBytes": "文件字节数", "sha256": "文件内容 SHA-256 摘要",
    "issuedAt": "证书签发时间", "certificateNo": "可公开验证的唯一证书编号",
    "amount": "本次积分变动值", "direction": "积分流入、流出或人工调整方向", "stock": "当前可兑换库存",
    "score": "服务端计算或模型返回的评分", "passed": "本次安全答题是否达到通过分数",
    "provider": "模型服务供应商标识", "model": "实际调用的模型名称", "errorMessage": "脱敏后的失败原因摘要",
    "conversationId": "关联聊天会话主键", "senderId": "消息发送者用户主键", "message": "聊天消息纯文本内容",
    "visibility": "知识或 Skill 的可见范围", "sensitivity": "内容敏感等级", "roleScope": "允许使用该内容的角色范围",
    "studentNo": "学生学号；未提供时允许为空", "college": "用户所属学院", "grade": "用户所在年级",
    "authorizationLevel": "账号当前授权等级", "pendingRegistrationApproval": "注册申请是否等待人工批准",
    "channelTitle": "通知所属频道的展示名称", "iconType": "通知状态图标类型", "menuScope": "通知对应的菜单未读范围",
    "entityType": "通知关联业务实体类型", "route": "点击通知后进入的平台内部路由",
    "emailStatus": "配套邮件的投递状态", "emailError": "邮件投递失败的脱敏错误摘要",
    "totalScore": "本次安全答题卷面总分", "threshold": "本次答题采用的通过分数",
    "type": "记录类型，由对应 Prisma 枚举约束", "stem": "安全题目题干", "options": "题目选项及多选答案结构",
    "explanation": "题目答案解释", "source": "记录的创建来源或能力来源", "active": "记录当前是否启用",
    "deviceFingerprint": "签署责任书时采集的脱敏设备指纹",
    "reportNo": "平台生成的唯一漏洞编号", "assetName": "受影响资产的展示名称或脱敏标识",
    "ownerUnit": "资产责任单位；尚未确认时允许为空", "vulnerabilityType": "漏洞类型名称",
    "selfSeverity": "提交者自评漏洞等级", "finalSeverity": "审核人员确认的最终漏洞等级",
    "reproductionSteps": "复现漏洞所需的步骤说明", "impactAssessment": "漏洞影响范围和风险分析",
    "repairSuggestion": "针对漏洞成因提出的修复建议", "requestPacket": "触发漏洞的脱敏请求包文本",
    "detailMarkdown": "正式漏洞报告的完整 Markdown 正文", "bodyMarkdown": "草稿阶段保存的 Markdown 正文",
    "testBoundary": "测试授权边界和禁止行为声明", "aiParsed": "报告内容是否经过 AI 辅助解析",
    "kind": "附件用途类型，包括报告、复测证据、正文图片和补充材料", "originalName": "用户上传文件的原始名称",
    "action": "业务事件或审计动作编码", "fromStatus": "状态迁移前的漏洞状态", "toStatus": "状态迁移后的漏洞状态",
    "deviceInfo": "操作发生时的脱敏设备信息", "decision": "审核人员作出的审核决定",
    "result": "复测、审计或投递执行结果", "evidenceSha256": "复测证据内容的 SHA-256 摘要",
    "verifyCode": "证书内部核验码", "publicVerify": "证书是否允许通过公开接口核验",
    "contribution": "证书中展示的安全贡献说明", "revokeReason": "撤销证书的原因",
    "templateSnapshot": "签发时固化的证书模板配置快照", "points": "本次积分变动的绝对值",
    "balanceAfter": "该笔流水提交后的用户积分余额", "category": "商品或知识内容分类",
    "pointsCost": "兑换商品需要扣除的积分", "perUserLimit": "单个用户允许兑换该商品的数量上限",
    "pickupLocation": "实物商品领取地点", "pickupInstructions": "实物商品领取说明",
    "sortOrder": "列表展示排序值", "pickupCode": "实物商品领取核销码",
    "modelName": "本次 AI Run 实际使用的模型名称", "purpose": "本次 AI 调用的业务用途",
    "policySnapshot": "运行创建时固化的模型策略快照", "inputTokens": "模型输入 Token 数量",
    "outputTokens": "模型输出 Token 数量", "errorCode": "AI Run 失败时的稳定错误编码",
    "aiSpace": "Skill 或 AI 能力所属的用户、审核员或管理员空间",
    "visibleToUsers": "系统 Skill 是否在用户侧可见", "defaultPinnedForUsers": "系统 Skill 是否默认固定到用户输入区",
    "slug": "Skill 的稳定短标识", "lifecycleStatus": "Skill 当前生命周期状态",
    "contentSha256": "Skill 当前内容包的 SHA-256 摘要", "installedPath": "系统 Skill 的受控安装路径",
    "syncError": "Skill 同步失败的脱敏错误信息", "canonicalQuestion": "知识条目归一化后的代表问题",
    "canonicalAnswer": "经审核的标准答案", "maintenanceStatus": "知识质量维护状态",
    "embeddingModel": "生成检索向量时使用的嵌入模型", "embeddingDimension": "知识向量维度",
    "generalityScore": "知识对同类问题的通用性评分", "safetyScore": "知识内容安全性评分",
    "privacyRiskScore": "知识包含隐私风险的评分", "educationalValueScore": "知识对学习场景的教育价值评分",
    "reuseValueScore": "知识后续复用价值评分", "answerQualityScore": "答案完整性和准确性的质量评分",
    "platformRelevanceScore": "知识与平台业务相关性的评分", "confidenceScore": "自动判断结果的置信评分",
    "autoApprovalReason": "满足自动入库规则的判定依据", "reviewNote": "管理员人工审核知识时填写的意见",
    "closedReason": "聊天会话被关闭的原因", "role": "用户角色或会话成员角色",
    "targetType": "审计动作指向的业务对象类型", "key": "系统设置的唯一配置键", "value": "系统设置的结构化配置值",
}


def field_meaning(model, name, field_type, attrs, model_names):
    if name in FIELD_MEANINGS:
        result = FIELD_MEANINGS[name]
    elif name.endswith("Id"):
        result = f"关联 {name[:-2]} 记录的主键"
    elif name.endswith("At"):
        result = f"{name[:-2]} 事件发生时间"
    elif name.endswith("Count"):
        result = f"{name[:-5]} 累计数量"
    elif name.endswith("Bytes"):
        result = f"{name[:-5]} 的字节容量"
    elif name.endswith("Hash"):
        result = f"{name[:-4]} 的不可逆摘要"
    elif name.endswith("Key"):
        result = f"{name[:-3]} 的稳定键值"
    elif name.endswith("Url") or name.endswith("URL"):
        result = f"{name} 对应的受控 URL"
    elif name.startswith("is") or name.startswith("has") or name.endswith("Enabled") or name.endswith("Required"):
        result = f"控制 {name} 条件是否成立的布尔标记"
    elif name.endswith("Json") or field_type.rstrip("?") == "Json":
        result = f"{name} 的结构化 JSON 数据"
    elif name.endswith("Text") or name.endswith("Summary"):
        result = f"{name} 对应的文本内容"
    else:
        result = f"{model} 模型中的 {name} 属性"
    qualifiers = []
    if "@unique" in attrs:
        qualifiers.append("值必须唯一")
    relation = re.search(r"@relation\([^)]*fields:\s*\[([^]]+)]", attrs)
    if relation:
        qualifiers.append(f"外键字段：{relation.group(1)}")
    if qualifiers:
        result += "；" + "；".join(qualifiers)
    return result


def prisma_to_physical(field_type, attrs, enum_names):
    optional = field_type.endswith("?")
    base = field_type.rstrip("?[]")
    db_type = {
        "String": "text", "Int": "integer", "BigInt": "bigint", "Boolean": "boolean",
        "DateTime": "timestamp", "Json": "jsonb", "Bytes": "bytea", "Float": "double precision", "Decimal": "decimal",
    }.get(base, base if base in enum_names else "relation")
    length = "—"
    varchar = re.search(r"@db\.VarChar\((\d+)\)", attrs)
    char = re.search(r"@db\.Char\((\d+)\)", attrs)
    decimal = re.search(r"@db\.Decimal\((\d+)\s*,\s*(\d+)\)", attrs)
    if varchar:
        db_type, length = "varchar", varchar.group(1)
    elif char:
        db_type, length = "char", char.group(1)
    elif decimal:
        db_type, length = "decimal", f"{decimal.group(1)},{decimal.group(2)}"
    elif db_type == "integer":
        length = "32 bit"
    elif db_type == "bigint":
        length = "64 bit"
    nullable = "可空" if optional else "非空"
    return db_type, length, nullable


def scalar_fields(model, models, enums):
    result = []
    model_names = set(models)
    enum_names = set(enums)
    for name, field_type, attrs in models[model]["fields"]:
        base = field_type.rstrip("?[]")
        if field_type.endswith("[]") or base in model_names:
            continue
        db_type, length, nullable = prisma_to_physical(field_type, attrs, enum_names)
        primary = "是" if "@id" in attrs else "否"
        description = field_meaning(model, name, field_type, attrs, model_names)
        constraints = [nullable]
        if "@unique" in attrs:
            constraints.append("唯一")
        if "@default" in attrs:
            constraints.append("有默认值")
        description += "；" + "、".join(constraints)
        result.append((str(len(result) + 1), name, db_type, length, primary, description))
    return result


def model_domain(name):
    groups = [
        ("身份与权限", ("User", "Role", "Permission", "Authorization", "Email", "Human", "Refresh", "Security", "Responsibility")),
        ("漏洞响应", ("Asset", "Vulnerability", "Review", "Retest")),
        ("证书与激励", ("Certificate", "Point", "Mall")),
        ("内容", ("Learning", "Announcement")),
        ("AI", ("Ai",)),
        ("聊天通知", ("Chat", "PlatformNotification")),
        ("平台治理", ("Audit", "School", "System")),
    ]
    for domain, prefixes in groups:
        if name.startswith(prefixes):
            return domain
    return "其他"


def add_cover(doc, reference):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run(PROJECT), "黑体", 22, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(16)
    set_run_font(p.add_run("系统设计说明书"), "黑体", 26, True)
    for _ in range(4):
        doc.add_paragraph()
    rows = [
        ("项  目  名  称", PROJECT), ("产　品　类　型", "Web 应用系统"), ("项  目  编  号.", ""),
        ("版    本    号", VERSION), ("状          态", "受控稿"), ("作          者", ""),
        ("所  属  部  门", ""), ("提  交  日  期", "2026年8月26日"),
    ]
    append_cloned_table(doc, reference, "cover", rows, has_header=False)
    page_break(doc)


def add_document_control(doc, reference):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("文档控制"), "黑体", 18, True)
    heading(doc, "文档更新记录", 2)
    add_template_table(doc, reference, "revision", "0-1", "文档更新记录", ["日期", "更新人", "版本", "备注"], [
        ("2026年8月26日", "", VERSION, "全面修订体系结构、模块协作、数据设计、安全控制、部署架构及接口追踪内容"),
    ])
    heading(doc, "文档审核记录", 2)
    add_template_table(doc, reference, "review", "0-2", "文档审核记录", ["日期", "审核人", "职务", "审核版本", "备注"], [
        ("", "", "", VERSION, ""),
    ])
    heading(doc, "文档发行范围", 2)
    add_template_table(doc, reference, "distribution", "0-3", "文档发行范围", ["抄送部门/个人", "说明"], [
        ("项目建设、测试、部署、运维及评审相关人员", "用于系统实现、验证、部署和维护"),
    ])
    page_break(doc)


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    set_run_font(p.add_run("目  录"), "黑体", 18, True)
    marker = doc.add_paragraph("[[DOC_TOC]]")
    marker.paragraph_format.first_line_indent = Pt(0)
    page_break(doc)


def chapter_1(doc, reference):
    heading(doc, "1 引言", 1)
    heading(doc, "1.1 编写目的", 2)
    body(doc, "本文档依据软件需求规格说明书，对校园漏洞响应与安全学习平台的总体架构、功能构件、业务流程、接口契约、数据模型、安全机制、部署结构及运行保障方案进行系统化设计说明。本文档作为系统建设、测试验证、部署实施、运行维护和变更控制的技术依据，并用于保持需求、设计与实现之间的可追溯性。")
    heading(doc, "1.2 预期读者和阅读建议", 2)
    bullets(doc, [
        "系统架构与开发人员重点阅读第3章至第5章，据此落实构件边界、接口契约、事务控制和数据约束。",
        "测试人员重点阅读第4章中的业务状态、异常路径和授权边界，并结合附件中的接口、权限和枚举清单建立验证矩阵。",
        "部署与运行维护人员重点阅读第3章的部署架构、安全控制、外部依赖、可观测性和故障降级设计。",
        "业务审核与管理人员重点阅读漏洞响应、通知协同、AI辅助、知识治理、Skill管理和审计追踪设计，确认人工决策责任边界。",
    ])
    heading(doc, "1.3 文档约定与设计依据", 2)
    body(doc, "本说明书采用需求约束、体系结构、详细设计、数据设计和运行设计逐层展开的组织方式。设计内容以经确认的软件需求、系统实现基线、数据模式及迁移记录、依赖与运行配置、部署配置和有效技术文档为依据；当不同资料存在差异时，以能够被系统行为、数据约束和配置项共同验证的设计事实为准。接口路径、权限编码、状态名称、数据模型和配置项采用等宽技术标识语义，但正文保持统一中文技术表述。")
    heading(doc, "1.4 参考资料", 2)
    add_template_table(doc, reference, "five", "1-1", "参考资料", ["编号", "资料名称", "版本/基线", "更新时间", "用途"], [
        ("1", "软件需求规格说明书", "现行受控版本", "2026-08-26", "业务范围、参与者、功能需求和非功能需求依据"),
        ("2", "系统实现与模块配置基线", VERSION, "2026-08-26", "构件职责、接口行为、授权规则和业务流程依据"),
        ("3", "数据库模式与迁移记录", "Prisma 6", "2026-08-26", "数据模型、枚举、约束、索引和结构演进依据"),
        ("4", "软件依赖与生产构建配置", VERSION, "2026-08-26", "技术栈、依赖版本、构建和资源优化依据"),
        ("5", "部署与运行配置基线", VERSION, "2026-08-26", "容器化部署、网络接入、数据服务和外部依赖依据"),
    ])
    heading(doc, "1.5 缩写术语", 2)
    add_template_table(doc, reference, "glossary", "1-2", "缩写术语", ["术语", "说明"], [
        ("SRC", "校园授权范围内的漏洞接收、审核、修复、复测、归档和激励协同机制。"),
        ("RBAC", "Role-Based Access Control，基于角色并结合用户权限覆盖和资源归属的访问控制。"),
        ("DTO", "Data Transfer Object，控制接口输入字段、类型、长度、格式和枚举范围。"),
        ("BYOK", "Bring Your Own Key，用户在平台策略允许时配置自己的模型供应商凭据。"),
        ("AI Run", "一次可追踪的 AI 编排运行，包含模型、步骤、工具、用量、状态和错误记录。"),
        ("Skill", "经过约束和人工审核的声明式 AI 指令能力，不等同于知识库条目。"),
        ("知识条目", "从已解决会话沉淀、审核并用于相似问题检索增强的可复用内容。"),
        ("S3", "对象存储兼容协议；平台以私有对象键管理漏洞附件和资源文件。"),
    ])


def chapter_2(doc, reference):
    heading(doc, "2 系统需求概述", 1)
    heading(doc, "2.1 软件需求概述", 2)
    body(doc, "平台采用角色驱动的校园漏洞响应闭环。用户完成安全准入后，以结构化基础字段和 Markdown 报告提交漏洞；审核员和管理员完成技术审核、补充材料、修复跟踪、复测与归档；证书、积分、学习资源、通知和聊天提供协同与激励；AI 仅在人工触发时提供报告质量、审核建议和知识增强，不获得最终业务决定权。")
    heading(doc, "2.2 项目功能需求列表", 2)
    add_template_table(doc, reference, "three", "2-1", "项目功能需求表", ["序号", "功能域", "实现范围"], [
        (str(index), name, scope) for index, (name, scope) in enumerate(FUNCTION_DOMAINS, 1)
    ])
    heading(doc, "2.3 项目非功能需求列表", 2)
    add_template_table(doc, reference, "nonfunctional", "2-2", "项目非功能需求表", ["序号", "质量属性", "设计约束"], [
        (str(index), name, detail) for index, (name, detail) in enumerate(NONFUNCTIONAL, 1)
    ])
    heading(doc, "2.4 项目用例与设计边界", 2)
    body(doc, "系统参与者包括未登录访问者、用户、审核员和管理员。外部依赖包括 SMTP 邮件、AI 模型供应商、S3 兼容对象存储和 ClamAV。外部系统只提供基础能力，不承担漏洞状态迁移和最终审核。")
    add_uml_placeholder(doc, "2-1", "系统总体用例图", "参与者：未登录访问者、用户、审核员、管理员。用例：公开访问、认证、准入、漏洞提交、审核、补充、复测、证书、积分、学习、通知、聊天、AI、知识库、Skill和平台治理。外部系统边界单独表示。")
    heading(doc, "2.5 设计追踪原则", 2)
    body(doc, "每个核心功能在第4章明确设计构件、接口契约、授权条件、状态变化、事务边界和异常处理；第5章说明逻辑数据域、物理模型、关系约束和并发控制；第6章规定界面结构、交互一致性、可访问性和性能约束。接口、权限、数据模型和枚举清单作为附件提供实现追踪依据。")


def chapter_3(doc, reference):
    heading(doc, "3 系统总体设计", 1)
    heading(doc, "3.1 运行环境说明", 2)
    add_template_table(doc, reference, "three", "3-1", "技术与运行环境", ["序号", "层次", "技术、版本与设计依据"], [
        (str(index), layer, f"{technology}；依据：{source}") for index, (layer, technology, source) in enumerate(TECHNOLOGY, 1)
    ])
    body(doc, "系统采用容器化部署方式，各应用与基础设施服务通过受控网络进行通信。部署设计将访问接入、静态资源服务、应用服务、关系数据库、缓存、对象存储和安全检测服务分层隔离，并通过环境配置管理连接信息、密钥和运行参数。计算资源规格、网络映射、持久化卷和备份策略属于部署配置项，应依据容量规划、性能测试和运行保障制度确定。")
    heading(doc, "3.2 系统部署与网络设计", 2)
    body(doc, "客户端通过 HTTPS 接入反向代理服务。静态资源请求由 Web 服务处理，REST 请求转发至应用服务，实时通信请求通过 WebSocket 升级机制转发至 Socket.IO 网关。应用服务访问 PostgreSQL、Redis、S3 兼容对象存储和恶意文件检测服务，并依据系统配置调用 SMTP 服务与 AI 模型供应商。除反向代理接入端口外，数据服务和内部基础设施端口不得直接暴露至公共网络。")
    add_uml_placeholder(doc, "3-1", "系统部署架构图", "UML 部署图。展示访问接入、容器化应用、数据服务、对象存储、安全检测和外部服务边界。")
    heading(doc, "3.3 软件系统架构设计", 2)
    add_template_table(doc, reference, "three", "3-2", "软件架构分层", ["序号", "层次", "核心构件与技术职责"], [
        ("1", "表现层", "React 页面、公共组件、路由守卫和服务封装；呈现四类端并处理交互状态，不替代服务端授权。"),
        ("2", "接入层", "Nginx Proxy Manager、Web Nginx、NestJS Main；负责 TLS、静态资源、API/WebSocket 转发、CORS、安全头和请求编号。"),
        ("3", "应用层", "NestJS Controller、Guard、Service、Worker、Gateway；负责输入、权限、状态机、事务、维护任务和实时通信。"),
        ("4", "数据访问层", "PrismaService；以 Prisma Schema 和 migration 管理 PostgreSQL 结构、关系、索引和事务。"),
        ("5", "基础设施层", "PostgreSQL、Redis、S3、ClamAV、SMTP、AI Provider；提供持久化、共享限流、文件检测、邮件和模型推理。"),
    ])
    add_uml_placeholder(doc, "3-2", "软件组件架构图", "UML 组件图。按表现层、接入层、应用层、数据访问层和基础设施层排列，标出 REST、Socket.IO、Prisma、Redis 和 S3 调用接口。")
    heading(doc, "3.4 前端架构设计", 2)
    body(doc, "前端入口 App.tsx 对管理员、审核员、用户和认证弹窗进行 lazy 导入；各角色 Home 再对业务页面二次懒加载。routeConfig.ts 统一菜单路由，apiClient.ts 统一携带身份、解析请求编号和映射错误。大型 Three.js 场景、非首屏模块和聊天依赖按需加载，Vite 将 React、Three.js、Motion 和实时通信依赖拆为独立生产块。")
    body(doc, "共享层负责主题、全局加载、全局提示、模态层级、响应式折叠、未读徽标和减少动态效果偏好。页面服务模块封装 REST 路径和类型，避免组件直接拼接认证逻辑。Socket.IO 在用户退出或页面销毁时解除监听，避免重复连接和事件订阅。")
    heading(doc, "3.5 后端模块与依赖设计", 2)
    body(doc, "AppModule 注册 Security、Prisma、ObjectStorage、Health、HumanVerification、Auth、Account、AI、AI Knowledge、Access、SecurityExam、Vulnerabilities、Attachments、Chat、Certificates、Announcements、AuditLogs、Learning、PointsMall、Notifications、SystemSettings、Users 和 Dashboard。模块化单体保持单次部署和跨域事务能力，同时以 NestJS 模块边界控制依赖。")
    add_uml_placeholder(doc, "3-3", "后端模块依赖图", "UML 包图。AppModule 位于顶层；业务模块按身份准入、漏洞响应、激励内容、通知沟通、AI和平台治理分组；Security、Prisma、ObjectStorage作为共享底层。")
    heading(doc, "3.6 数据与存储设计", 2)
    body(doc, "PostgreSQL 保存业务主数据、状态、关系和审计记录，Redis 保存跨实例共享限流与短期运行状态，对象存储保存原始附件、学习资源、证书文件和外观资源。关系数据库仅保存私有对象键及其元数据，文件访问必须经过应用服务的身份、权限、资源归属和业务状态校验。数据库模式迁移构成结构演进的唯一受控通道。")
    add_uml_placeholder(doc, "3-4", "核心领域模型图", "UML 类图。展示 User、VulnerabilityReport、Attachment、Supplement、Review、Retest、Certificate、PointLedger、Notification、Chat、AiAgentRun、AiKnowledgeEntry和AiSkill及其基数。")
    heading(doc, "3.7 安全架构设计", 2)
    body(doc, "应用启动层统一设置 CSP、CORP、Permissions-Policy、Referrer-Policy、X-Content-Type-Options、X-Frame-Options 和生产环境 HSTS；跨域访问仅允许受控来源列表。可信代理配置用于解析真实客户端地址，外部传入的 X-Forwarded-For 不直接作为审计依据，系统以受信任代理链解析后的客户端地址覆盖不可信输入，降低审计来源伪造风险。")
    body(doc, "JWT、RBAC、用户权限覆盖、资源归属和状态前置条件共同控制访问。JSON 请求体限制为 5 MB，文件接口另设类型、大小、数量、速率和恶意内容检测。模型密钥以独立 32 字节密钥加密，出站 URL 防止 SSRF，AI 出站内容按字段脱敏。")
    add_uml_placeholder(doc, "3-5", "安全控制架构图", "安全工程图。展示 TLS/反向代理、安全头、JWT、RBAC、资源归属、DTO、限流、ClamAV、对象存储授权、密钥加密、AI脱敏、审计和备份控制。")
    heading(doc, "3.8 事务、可观测性与降级", 2)
    body(doc, "关键写操作使用 Prisma $transaction 保证状态、事件、流水、通知记录和审计一致。外部邮件、实时推送和 AI 供应商调用置于可降级边界：核心事务提交后再派发，失败保留状态并写入可追踪记录。X-Request-Id 关联前端错误、HTTP 结构化日志和审计；默认只记录错误请求或耗时超过 1000 ms 的请求，可通过 HTTP_LOG_ALL 扩展。")
    add_template_table(doc, reference, "three", "3-3", "外部依赖与降级策略", ["序号", "依赖", "用途与不可用时行为"], [
        ("1", "PostgreSQL", "核心业务数据与事务。健康检查失败时停止受影响业务，禁止静默降级为内存数据。"),
        ("2", "Redis", "生产共享限流和短期状态。生产环境为必需配置，缺少 REDIS_URL 时启动失败。"),
        ("3", "S3 兼容对象存储", "保存附件、资源、证书和外观文件。不可用时拒绝相关写入或读取，避免无效对象引用。"),
        ("4", "恶意文件检测服务", "检测上传内容。检测服务不可用或无法形成可信结论时，按安全策略拒绝高风险上传并记录异常。"),
        ("5", "SMTP", "发送验证码与重要邮件。站内业务继续；失败不自动重试，写入通知投递审计。"),
        ("6", "AI 模型供应商", "提供问答、分析和审核建议。失败时返回统一错误，不改变漏洞状态，不阻断人工流程。"),
    ])


def chapter_4(doc, reference, endpoints):
    heading(doc, "4 功能模块详细设计", 1)
    body(doc, "本章按照业务能力划分详细说明各功能模块的构件职责、处理流程、状态约束、接口契约、事务边界、异常处理和审计要求。接口统一采用 /api 路径前缀；访问控制同时包含角色权限、用户权限覆盖、资源归属和业务状态前置条件。主要接口用于说明模块对外契约，完整接口清单见附件A。")
    for module_index, module in enumerate(MODULES, 1):
        heading(doc, f"{module['number']} {module['title']}", 2)
        heading(doc, f"{module['number']}.1 设计目标与构件职责", 3)
        body(doc, f"{module['title']}模块用于落实相应业务流程及其安全、数据和审计约束。模块通过接口层接收请求，由领域服务执行规则和状态控制，并通过数据访问与基础设施构件完成持久化及外部协作。各构件职责如表4-{module_index}-1所示。")
        add_template_table(doc, reference, "three", f"4-{module_index}-1", f"{module['title']}构件职责", ["序号", "设计构件", "职责说明"], [
            (str(index), component, f"{layer}：{responsibility}")
            for index, (layer, component, responsibility) in enumerate(module["components"], 1)
        ])
        heading(doc, f"{module['number']}.2 技术实现", 3)
        for paragraph in module["design"]:
            body(doc, paragraph)
        heading(doc, f"{module['number']}.3 事务、异常与审计", 3)
        body(doc, module["transactions"])
        keywords = set()
        for file_path in module["files"]:
            name = Path(file_path).parts[-2] if "/" in file_path else file_path
            keywords.add(name)
        module_endpoints = [endpoint for endpoint in endpoints if any(f"/modules/{keyword}/" in endpoint["source"] for keyword in keywords)]
        module_endpoints = module_endpoints[:6]
        if module_endpoints:
            heading(doc, f"{module['number']}.4 主要接口", 3)
            add_template_table(doc, reference, "three", f"4-{module_index}-2", f"{module['title']}主要接口", ["序号", "方法与路径", "处理与授权约束"], [
                (
                    str(index),
                    f"{endpoint['method']} {endpoint['path']}",
                    f"处理器：{endpoint['handler']}；授权要求：{endpoint['permissions']}。",
                )
                for index, endpoint in enumerate(module_endpoints, 1)
            ])
        add_uml_placeholder(doc, f"4-{module_index}", f"{module['title']}工程图", module["uml"])


def chapter_5(doc, reference, enums, models):
    heading(doc, "5 数据库设计说明", 1)
    body(doc, "系统使用 PostgreSQL 作为关系数据库，并通过 Prisma Schema 统一定义物理模型、字段类型、枚举、关系、唯一约束和索引。当前设计包含76个数据模型和46个枚举类型；枚举映射为数据库枚举，实体关联通过外键、唯一约束和组合索引实现。")
    heading(doc, "5.1 数据域划分", 2)
    domains = {}
    for name in models:
        domains.setdefault(model_domain(name), []).append(name)
    add_template_table(doc, reference, "three", "5-1", "数据域与模型分布", ["序号", "数据域", "模型数量与名称"], [
        (str(index), domain, f"共 {len(names)} 个：{'、'.join(names)}") for index, (domain, names) in enumerate(domains.items(), 1)
    ])
    heading(doc, "5.2 关系、约束与迁移", 2)
    body(doc, "实体主键以全局唯一字符串标识为主；业务编号、证书编号、权限编码和事件键采用唯一约束。高频查询依据用户、状态、时间和业务对象建立组合索引。删除策略根据实体生命周期分别采用级联删除、外键置空或历史记录保留。数据库结构变更必须通过受控迁移脚本发布，禁止绕过迁移流程直接修改生产结构。")
    body(doc, "知识库与声明式 Skill 属于相互独立的数据域。历史兼容枚举仅用于识别既有数据，不构成现行业务转换路径；任何知识条目不得通过状态变更直接转换为 Skill。")
    add_uml_placeholder(doc, "5-1", "数据库核心实体关系图", "ER/UML 类图。只展开核心实体和基数，辅助运行步骤、版本、反馈、维护任务和审计表作为分组附注。")
    heading(doc, "5.3 核心物理表设计", 2)
    for index, model in enumerate(CORE_MODELS, 1):
        heading(doc, f"5.3.{index} {model} 表", 3)
        rows = scalar_fields(model, models, enums)
        add_template_table(doc, reference, "physical", f"5-{index + 1}", f"{model} 物理字段设计", ["序号", "列名", "数据类型", "长度", "主键", "说明"], rows)
        directives = models[model]["directives"]
        relations = []
        for name, field_type, attrs in models[model]["fields"]:
            base = field_type.rstrip("?[]")
            if base in models:
                cardinality = "一对多集合" if field_type.endswith("[]") else "可选关联" if field_type.endswith("?") else "必选关联"
                relations.append(f"{name}:{base}（{cardinality}）")
        if relations:
            body(doc, "关系设计：" + "；".join(relations) + "。")
        if directives:
            body(doc, "模型级约束与索引：" + "；".join(directives) + "。")
    heading(doc, "5.4 一致性与并发控制", 2)
    bullets(doc, [
        "漏洞草稿使用 version 条件更新实现乐观并发，冲突返回 409，不静默覆盖。",
        "审核、复测、证书和状态事件在事务中按当前状态校验，阻止越级和重复迁移。",
        "商城兑换以条件库存更新、兑换状态和积分流水保证不超卖、不重复扣分。",
        "通知记录在核心事务内创建，邮件和 WebSocket 在事务提交后派发。",
        "AI Run 使用运行状态和条件更新处理取消、完成和失败，避免结束状态被覆盖。",
    ])
    heading(doc, "5.5 数据安全、备份与保留", 2)
    body(doc, "密码只保存派生摘要；模型凭据以 AI_CREDENTIAL_ENCRYPTION_KEY 加密；JWT 密钥不入库。附件文件保存于私有对象存储，数据库保存对象键、摘要、类型、大小和授权关系。审计和日志不得记录完整密码、令牌、模型密钥或未经处理的漏洞原文。")
    body(doc, "备份范围必须同时覆盖 PostgreSQL 业务数据和对象存储文件，并通过定期恢复演练验证备份的完整性与可恢复性。备份周期、保留期限、恢复点目标和恢复时间目标应纳入运行维护制度，并依据业务连续性要求、数据增长量和基础设施能力进行配置。")


def chapter_6(doc, reference):
    heading(doc, "6 软件界面设计", 1)
    heading(doc, "6.1 公共界面架构", 2)
    body(doc, "界面以任务和数据为中心，统一使用响应式内容宽度、公共导航、全局加载、全局提示、模态层级和明暗主题。图标只用于明确操作或状态，不作为无引导意义的装饰。页面组件不承担最终权限判断，接口返回和业务守卫决定可执行能力。")
    add_template_table(doc, reference, "three", "6-1", "界面公共设计约束", ["序号", "设计项", "实现机制与验收重点"], [
        ("1", "响应式", "稳定内容宽度、断点折叠侧栏、可换行工具栏和受控弹窗尺寸；窄屏不遮挡正文。"),
        ("2", "主题", "CSS 变量、Tailwind dark 变体和系统外观设置；输入、卡片、遮罩、图表和 AI 回复均可读。"),
        ("3", "加载", "全局圆形加载、字体与首屏资源就绪检测和淡出；刷新时不暴露错误字体或未完成首屏。"),
        ("4", "提示", "统一全局提示服务；成功、错误、警告和处理中样式一致，错误保留请求编号。"),
        ("5", "模态", "统一门户层和遮罩层级；详情或大图打开时菜单不可操作，空白处和关闭键可退出。"),
        ("6", "动效", "Motion/CSS 过渡、reduced-motion 和设备能力降级；低配设备不持续消耗资源。"),
        ("7", "输入", "前端约束与服务端 DTO 双重校验；长度、格式、枚举、必填和上传限制明确。"),
    ])
    for index, (number, title, route, role, scope, masking) in enumerate(UI_SECTIONS, 1):
        heading(doc, f"{number} {title}", 2)
        body(doc, f"访问路径：{route}。目标角色：{role}。界面截图应覆盖：{scope}。截图用于说明界面结构和交互状态，不作为权限有效性或业务正确性的唯一证明；相关验收仍以接口响应、状态变化和审计记录为准。")
        add_screenshot_placeholder(doc, f"6-{index}", f"{title}界面图", f"页面/路由：{route}\n目标角色：{role}\n建议范围：{scope}\n脱敏要求：{masking}")
    heading(doc, "6.10 可访问性与性能", 2)
    body(doc, "交互控件使用语义化 button、input、dialog 和 aria 属性；键盘焦点在弹窗打开和关闭时正确迁移；颜色不作为唯一状态表达。字体不按视口宽度缩放，固定格式组件使用稳定尺寸。非首屏页面、Three.js 和大依赖按需加载，不可见场景停止持续渲染并释放 WebGL 资源。")


def attachments(doc, reference, endpoints, enums, models):
    heading(doc, "附件", 1)
    heading(doc, "附件A 接口清单", 2)
    body(doc, "本清单按接口控制器聚合系统 REST 接口，记录请求方法、完整路径、处理器和可直接识别的权限要求。未在方法级标注权限编码的接口，仍可能受到类级身份守卫、角色规则、资源归属和业务状态条件约束；接口清单不得替代详细授权设计。")
    grouped_endpoints = {}
    for endpoint in endpoints:
        grouped_endpoints.setdefault(endpoint["source"], []).append(endpoint)
    endpoint_rows = []
    for source, entries in grouped_endpoints.items():
        details = "\n".join(f"{entry['method']} {entry['path']} [{entry['handler']}]" for entry in entries)
        permissions = sorted({entry["permissions"] for entry in entries})
        details = f"授权要求：{'；'.join(permissions)}\n{details}"
        controller_info = f"{Path(source).stem}\n接口数：{len(entries)}"
        endpoint_rows.append((controller_info, details))
    add_template_table(
        doc,
        reference,
        "glossary",
        "附表A-1",
        "REST 接口清单",
        ["接口控制器", "接口与授权明细"],
        endpoint_rows,
        column_widths=(0.34, 0.66),
    )
    heading(doc, "附件B 权限清单", 2)
    permission_source = (ROOT / "backend/src/modules/rbac/permissions.ts").read_text(encoding="utf-8")
    definitions = re.findall(r"\['([^']+)',\s*'([^']+)'\]", permission_source)
    role_rows = []
    for index, (code, label) in enumerate(definitions, 1):
        roles = []
        for role in ("STUDENT", "AUDITOR", "ADMIN"):
            block = re.search(rf"{role}:\s*\[(.*?)]", permission_source, re.S)
            if block and f"'{code}'" in block.group(1):
                roles.append(role)
        role_rows.append((str(index), code, f"{label}；角色基线：{'、'.join(roles) or '按用户权限覆盖授予'}"))
    add_template_table(doc, reference, "three", "附表B-1", "RBAC 权限定义", ["序号", "权限编码", "权限说明与角色基线"], role_rows)
    heading(doc, "附件C Prisma 模型清单", 2)
    add_template_table(doc, reference, "three", "附表C-1", "Prisma 模型清单", ["序号", "数据域", "模型名称"], [
        (str(index), model_domain(name), name) for index, name in enumerate(models, 1)
    ])
    heading(doc, "附件D 状态与枚举清单", 2)
    add_template_table(doc, reference, "three", "附表D-1", "Prisma 枚举清单", ["序号", "枚举名称", "枚举值"], [
        (str(index), name, "、".join(values)) for index, (name, values) in enumerate(enums.items(), 1)
    ])
    body(doc, "说明：AiKnowledgeStatus.PROMOTED_TO_SKILL、AiKnowledgeSource.SKILL_PROMOTION、AiKnowledgeConversationStatus.SKILL_CANDIDATE 和 AiProblemSessionStatus.SKILL_CANDIDATE 为历史兼容值。现行迁移已解除知识库与 Skill 的活动关联。")
    heading(doc, "附件E UML 工程图清单", 2)
    add_template_table(doc, reference, "three", "附表E-1", "UML 工程图清单", ["图号", "图名/类型", "表达内容"], [
        (number, f"{title}（{kind}）", detail) for number, title, kind, detail in UML_LIST
    ])
    heading(doc, "附件F 界面截图清单", 2)
    add_template_table(doc, reference, "three", "附表F-1", "界面截图清单", ["图号", "页面/角色", "截图与脱敏要求"], [
        (f"图6-{index}", f"{title}；{role}", f"路由：{route}；范围：{scope}；脱敏：{masking}")
        for index, (_, title, route, role, scope, masking) in enumerate(UI_SECTIONS, 1)
    ])
    heading(doc, "附件G 运行参数配置项", 2)
    add_template_table(doc, reference, "three", "附表G-1", "运行参数配置项", ["序号", "参数", "属性与确定依据"], [
        (str(index), parameter, f"{status}；{method}") for index, (parameter, status, method) in enumerate(UNCONFIRMED, 1)
    ])


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def deduplicate(path):
    with ZipFile(path, "r") as source:
        names = list(dict.fromkeys(source.namelist()))
        payloads = {name: source.read(name) for name in names}
    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(temporary, "w", ZIP_DEFLATED) as target:
            for name in names:
                target.writestr(name, payloads[name])
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def main():
    reference = Document(REFERENCE)
    doc = Document(REFERENCE)
    remove_body(doc)
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        clear_header_footer(section)
    add_cover(doc, reference)
    add_document_control(doc, reference)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(body_section)
    add_header_footer(body_section)
    add_toc(doc)
    enums, models = parse_schema()
    endpoints = parse_endpoints()
    if len(models) != 76:
        raise RuntimeError(f"expected 76 Prisma models, found {len(models)}")
    chapter_1(doc, reference)
    chapter_2(doc, reference)
    chapter_3(doc, reference)
    chapter_4(doc, reference, endpoints)
    chapter_5(doc, reference, enums, models)
    chapter_6(doc, reference)
    attachments(doc, reference, endpoints, enums, models)
    set_update_fields(doc)
    doc.core_properties.title = f"{PROJECT}系统设计说明书"
    doc.core_properties.subject = "软件系统设计"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "校园漏洞响应, 安全学习, 系统设计, V2.10"
    doc.core_properties.comments = "系统设计受控文档。"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    deduplicate(OUTPUT)
    print(OUTPUT)
    print(f"models={len(models)} enums={len(enums)} endpoints={len(endpoints)} tables={len(doc.tables)} paragraphs={len(doc.paragraphs)}")


if __name__ == "__main__":
    main()
