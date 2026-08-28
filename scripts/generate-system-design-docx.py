#!/usr/bin/env python3
"""Generate the formal system design specification from the preserved template."""

from __future__ import annotations

from datetime import date
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from assets.system_design_content import (
    ARCH_COMPONENTS,
    CORE_MODELS,
    FIGURE_PLACEHOLDERS,
    FUNCTIONS,
    MODEL_GROUPS,
    MODULE_DESIGNS,
    NONFUNCTIONAL,
    ROLES,
    TECH_BASELINE,
    UI_AREAS,
)


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "scripts" / "assets" / "system-design-reference.docx"
OUTPUT = ROOT / "docs" / "校园漏洞响应与安全学习平台项目系统设计说明书.docx"
SCHEMA = ROOT / "database" / "prisma" / "schema.prisma"
PROJECT = "校园漏洞响应与安全学习平台"
VERSION = "V1.0"
GENERATED_ON = date(2026, 8, 26)


def set_run_font(run, name="宋体", size=12, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def remove_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def style_named(doc, name):
    for style in doc.styles:
        if style.name == name:
            return style
    raise KeyError(f"no style with name {name!r}")


def configure_styles(doc):
    normal = style_named(doc, "Normal")
    normal.font.name = "宋体"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = style_named(doc, name)
        style.font.name = "黑体"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style._element.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)


def configure_page(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)


def clear_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for container in (section.header, section.footer):
        for paragraph in list(container.paragraphs):
            paragraph._element.getparent().remove(paragraph._element)


def add_field(paragraph, instruction, fallback=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hp.add_run(f"{PROJECT}项目系统设计说明书"), "宋体", 9)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("第 "), "宋体", 9)
    add_field(fp, "PAGE", "1")
    set_run_font(fp.add_run(" 页"), "宋体", 9)
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        section._sectPr.append(pg)
    pg.set(qn("w:start"), "1")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def add_table_caption(doc, number, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(f"表 {number}  {title}"), "宋体", 10.5)


def add_table(doc, number, title, headers, rows, widths=None):
    add_table_caption(doc, number, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "D9E2F3")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        set_run_font(p.add_run(str(header)), "宋体", 10.5, True)
    for values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            set_run_font(p.add_run(str(value)), "宋体", 10.5)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
    return table


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), "宋体", 12, True)
        set_run_font(p.add_run(text[len(bold_prefix):]), "宋体", 12)
    else:
        set_run_font(p.add_run(text), "宋体", 12)
    return p


def add_bullets(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        set_run_font(p.add_run(f"（{index}）{item}"), "宋体", 12)


def add_figure_placeholder(doc, number, title, description):
    separator = doc.add_paragraph()
    separator.paragraph_format.space_before = Pt(0)
    separator.paragraph_format.space_after = Pt(0)
    separator.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, top=260, start=260, bottom=260, end=260)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    set_run_font(p.add_run("【图示占位】\n"), "黑体", 11, True, "666666")
    set_run_font(p.add_run(description), "宋体", 10.5, False, "555555")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = False
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    set_run_font(caption.add_run(f"图 {number}  {title}"), "宋体", 10.5)


def heading(doc, text, level):
    paragraph = doc.add_paragraph()
    paragraph.style = style_named(doc, f"Heading {level}")
    set_run_font(paragraph.add_run(text), "黑体", {1: 16, 2: 14, 3: 12}[level], True)
    return paragraph


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(PROJECT), "黑体", 22, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    set_run_font(p.add_run("项目系统设计说明书"), "黑体", 26, True)
    for _ in range(5):
        doc.add_paragraph()
    rows = [
        ("文档版本", VERSION),
        ("编制部门", "蓝山工作室运维安全部安全组"),
        ("编制日期", GENERATED_ON.strftime("%Y年%m月%d日")),
        ("文档状态", "正式版"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for label, value in rows:
        cells = table.add_row().cells
        for idx, text in enumerate((label, value)):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            set_run_font(p.add_run(text), "宋体", 12, idx == 0)
            set_cell_margins(cells[idx], 130, 130, 130, 130)
        set_cell_shading(cells[0], "D9E2F3")
    page_break(doc)


def add_control_pages(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("文档修订记录"), "黑体", 16, True)
    add_table(doc, "0-1", "文档修订记录", ["版本", "日期", "修订内容", "修订人"], [
        (VERSION, GENERATED_ON.strftime("%Y-%m-%d"), "依据当前系统实现重建系统设计说明书", "蓝山工作室运维安全部安全组")
    ], [2.0, 3.0, 7.0, 4.0])
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("文档审核记录"), "黑体", 16, True)
    add_table(doc, "0-2", "文档审核记录", ["审核轮次", "审核人", "审核日期", "审核意见", "处理结果"], [
        ("1", "", "", "", "")
    ], [1.8, 2.8, 2.8, 6.0, 2.6])
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("文档分发记录"), "黑体", 16, True)
    add_table(doc, "0-3", "文档分发记录", ["接收对象", "用途"], [
        ("项目建设、测试、运维及评审相关人员", "系统设计、实现、测试、部署和维护依据")
    ], [7.5, 8.5])


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目  录"), "黑体", 18, True)
    marker = doc.add_paragraph("[[DOC_TOC]]")
    marker.paragraph_format.first_line_indent = Pt(0)


def parse_prisma_models():
    source = SCHEMA.read_text(encoding="utf-8")
    models = {}
    for match in re.finditer(r"^model\s+(\w+)\s*\{(.*?)^\}", source, re.M | re.S):
        name, body = match.groups()
        fields = []
        indexes = []
        for raw in body.splitlines():
            line = raw.strip()
            if not line or line.startswith("//"):
                continue
            if line.startswith("@@"):
                indexes.append(line)
                continue
            if line.startswith("@"):  # defensive; model directives are handled above
                continue
            parts = line.split()
            if len(parts) < 2:
                continue
            field, field_type = parts[0], parts[1]
            attrs = " ".join(parts[2:])
            fields.append((field, field_type, attrs))
        models[name] = {"fields": fields, "indexes": indexes}
    return models


def set_update_fields_on_open(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def normalize_table_paragraphs(doc):
    for table in doc.tables:
        set_table_borders(table)
        for paragraph in table._tbl.iter(qn("w:p")):
            p_pr = paragraph.find(qn("w:pPr"))
            if p_pr is None:
                p_pr = OxmlElement("w:pPr")
                paragraph.insert(0, p_pr)
            numbering = p_pr.find(qn("w:numPr"))
            if numbering is not None:
                p_pr.remove(numbering)
            indentation = p_pr.find(qn("w:ind"))
            if indentation is None:
                indentation = OxmlElement("w:ind")
                p_pr.append(indentation)
            for attr in ("left", "right", "start", "end", "firstLine", "firstLineChars"):
                indentation.set(qn(f"w:{attr}"), "0")
            for attr in ("hanging", "hangingChars"):
                indentation.attrib.pop(qn(f"w:{attr}"), None)


def deduplicate_package(path):
    with ZipFile(path, "r") as source:
        names = list(dict.fromkeys(source.namelist()))
        payloads = {name: source.read(name) for name in names}
    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with ZipFile(temporary, "w", ZIP_DEFLATED) as target:
            for name in names:
                target.writestr(name, payloads[name])
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


FIELD_DESCRIPTIONS = {
    "id": "主键标识", "userId": "关联用户标识", "createdAt": "创建时间", "updatedAt": "更新时间",
    "deletedAt": "软删除时间", "status": "业务状态", "title": "标题", "description": "说明内容",
    "name": "名称", "email": "邮箱地址", "username": "登录账号", "passwordHash": "密码摘要",
    "roleId": "关联角色标识", "reportId": "关联漏洞报告标识", "vulnerabilityId": "关联漏洞标识",
    "content": "正文内容", "metadata": "扩展元数据", "version": "版本号", "type": "类型",
    "score": "评分", "passed": "是否通过", "expiresAt": "过期时间", "revokedAt": "撤销时间",
    "readAt": "读取时间", "issuedAt": "签发时间", "amount": "数量或变动值", "reason": "原因",
    "fileName": "原始文件名", "objectKey": "对象存储键", "mimeType": "媒体类型", "size": "文件大小",
    "provider": "模型供应商", "model": "模型标识", "errorMessage": "错误摘要",
}


def field_description(name, field_type):
    if name in FIELD_DESCRIPTIONS:
        return FIELD_DESCRIPTIONS[name]
    related_type = field_type.rstrip("?[]")
    if field_type.endswith("[]"):
        return f"关联的 {related_type} 记录集合"
    if related_type[:1].isupper() and related_type not in {"String", "Int", "BigInt", "Boolean", "DateTime", "Json", "Bytes", "Decimal", "Float"}:
        return f"关联的 {related_type} 记录"
    if name.endswith("Id"):
        return f"关联 {name[:-2]} 的标识"
    if name.startswith("is") or name.startswith("has"):
        return "业务布尔标记"
    if name.endswith("At"):
        return "业务发生时间"
    return f"{name} 业务字段"


def constraint_text(field_type, attrs):
    constraints = []
    if "@id" in attrs:
        constraints.append("主键")
    if "@unique" in attrs:
        constraints.append("唯一")
    if "@default" in attrs:
        default = re.search(r"@default\(([^)]*(?:\)[^)]*)?)\)", attrs)
        constraints.append(f"默认值 {default.group(1) if default else '已定义'}")
    if "@relation" in attrs:
        constraints.append("外键关系")
    if field_type.endswith("?"):
        constraints.append("可空")
    elif not field_type.endswith("[]"):
        constraints.append("非空")
    if field_type.endswith("[]"):
        constraints.append("关联集合")
    return "；".join(constraints) or "—"


def add_chapter_1(doc):
    heading(doc, "1 引言", 1)
    heading(doc, "1.1 编写目的", 2)
    add_body(doc, "本文档规定校园漏洞响应与安全学习平台的软件总体结构、模块边界、数据设计、接口协作、安全控制、部署方式和界面设计原则，作为开发、测试、部署、运维及后续变更评审的共同技术依据。")
    heading(doc, "1.2 项目范围", 2)
    add_body(doc, "平台面向校园授权范围内的漏洞响应与安全学习活动，覆盖安全准入、漏洞提交、审核复测、证书与积分、学习内容、通知沟通、AI 辅助及平台治理。平台不替代人工漏洞定性、风险评级和最终审核决策。")
    heading(doc, "1.3 读者对象", 2)
    add_body(doc, "本文档适用于产品与项目负责人、前后端开发人员、测试人员、部署运维人员、安全审核人员、数据库维护人员及经授权的评审人员。")
    heading(doc, "1.4 设计依据", 2)
    add_table(doc, "1-1", "设计依据", ["类别", "依据", "作用"], [
        ("需求基线", "校园漏洞响应与安全学习平台软件需求规格说明书", "确定系统边界、角色、流程和质量要求"),
        ("实现基线", "当前前端路由、NestJS 模块、控制器、服务及权限映射", "确认可实现的模块职责和接口关系"),
        ("数据基线", "database/prisma/schema.prisma 及迁移文件", "确认实体、约束和关系"),
        ("部署基线", "Docker 配置、服务器说明和更新脚本", "确认当前测试环境部署结构"),
        ("安全基线", "鉴权、RBAC、限流、审计、文件检测和密钥校验实现", "确定安全设计约束"),
    ], [2.6, 7.2, 6.2])
    heading(doc, "1.5 术语与缩略语", 2)
    add_table(doc, "1-2", "术语与缩略语", ["术语", "说明"], [
        ("SRC", "漏洞响应与协同处理机制，本平台用于校园授权范围内的漏洞报告与闭环管理。"),
        ("安全准入", "用户完成资料、邮箱验证、责任书签署和安全知识答题后获得受保护功能访问资格。"),
        ("RBAC", "基于角色的访问控制，并结合权限点和资源归属执行服务端授权。"),
        ("Skill", "经审核的声明式 AI 指令能力，与知识库条目相互独立。"),
        ("AI 知识库", "从已解决会话沉淀并经规则或人工审核的可复用问答知识集合。"),
        ("S3 兼容对象存储", "通过对象键保存漏洞附件、学习资源等文件的存储服务。"),
    ], [4.0, 12.0])
    heading(doc, "1.6 文档约定", 2)
    add_body(doc, "本文档中“必须”表示实现和验收所需约束，“应”表示在无明确冲突时遵循的设计要求，“可”表示允许的扩展。当前阶段暂不绘制工程图，图位以灰色占位框明确图型、节点和关系；后续替换时保持图号和图题不变。所有表题置于表上方，所有图题置于图下方。")


def add_chapter_2(doc):
    heading(doc, "2 项目需求分析", 1)
    heading(doc, "2.1 系统边界与参与者", 2)
    add_body(doc, "系统边界由公开访问域、登录用户业务域、审核业务域、管理治理域和外部基础服务构成。业务权限由后端统一判定，前端菜单和按钮仅作为交互层约束。")
    add_table(doc, "2-1", "参与者及职责", ["参与者", "主要职责"], ROLES, [4.0, 12.0])
    heading(doc, "2.2 功能需求映射", 2)
    add_table(doc, "2-2", "功能域与实现模块映射", ["功能域", "主要能力", "实现模块或路由"], FUNCTIONS, [3.0, 7.3, 5.7])
    heading(doc, "2.3 非功能需求约束", 2)
    add_table(doc, "2-3", "非功能设计约束", ["质量属性", "设计约束"], NONFUNCTIONAL, [3.2, 12.8])
    heading(doc, "2.4 关键业务规则", 2)
    add_bullets(doc, [
        "用户未完成安全准入时，不得提交正式漏洞或执行受保护操作。",
        "漏洞状态只能按服务端定义的合法迁移推进，AI 建议不直接改变状态。",
        "站内重要通知不可由用户关闭；已读后菜单红点应同步消失。",
        "用户间首次会话在对方回复前最多发送三条消息，管理员与审核员会话同样适用。",
        "只有管理员可创建和管理群聊；普通群人数上限为 200，全体群按平台人员范围维护。",
        "AI 调用按需触发；用户报告和审核建议失败时不得阻断核心漏洞流程。",
        "知识库条目不会创建、转换或安装为 Skill；用户 Skill 的公开申请全部人工审核。",
    ])
    heading(doc, "2.5 需求到设计的可追踪性", 2)
    rows = []
    for index, item in enumerate(FUNCTIONS, 1):
        rows.append((f"FR-{index:02d}", item[0], f"第 4 章相关模块；第 5 章相关数据表；第 6 章对应界面"))
    add_table(doc, "2-4", "需求设计追踪矩阵", ["需求编号", "需求域", "设计落点"], rows, [2.5, 4.2, 9.3])


def add_chapter_3(doc):
    heading(doc, "3 项目总体设计", 1)
    heading(doc, "3.1 设计目标与原则", 2)
    add_body(doc, "总体设计遵循业务闭环优先、服务端可信、领域模块隔离、关键动作可审计、外部能力可降级、数据最小披露和前端渐进增强原则。实现保持前后端分离，实时通信与普通接口共享身份和权限边界。")
    heading(doc, "3.2 技术与运行环境", 2)
    add_table(doc, "3-1", "技术基线", ["层次", "技术与版本"], TECH_BASELINE, [4.0, 12.0])
    add_body(doc, "测试环境已确认采用 Ubuntu 22.04、Docker 26 和 Nginx Proxy Manager。生产环境硬件规格、容灾级别及正式容量指标未确认，部署前应根据实际并发、附件规模和模型调用量完成容量评估。")
    heading(doc, "3.3 部署结构", 2)
    add_body(doc, "浏览器通过 HTTPS 访问反向代理。静态前端由 Web/Nginx 容器提供，API 与 Socket.IO 请求转发至 NestJS 后端。后端访问 PostgreSQL、Redis 和 S3 兼容对象存储，并按配置连接邮件与模型供应商。对象存储中的原始漏洞附件不经过 AI 脱敏改写。")
    num, title, desc = FIGURE_PLACEHOLDERS[0]
    add_figure_placeholder(doc, num, title, desc)
    heading(doc, "3.4 总体软件架构", 2)
    add_table(doc, "3-2", "总体架构分层", ["层次", "设计说明"], ARCH_COMPONENTS, [3.2, 12.8])
    num, title, desc = FIGURE_PLACEHOLDERS[1]
    add_figure_placeholder(doc, num, title, desc)
    heading(doc, "3.5 前端架构设计", 2)
    add_body(doc, "前端以 React 组件树和 React Router 路由为骨架，按公开端、用户端、审核员端、管理员端划分页面域。认证状态、主题、全局加载和统一提示由共享设施提供；大型页面、Three.js 场景和非首屏功能采用动态导入。接口层集中处理令牌、请求编号和错误映射，避免页面直接拼接底层协议。")
    add_bullets(doc, [
        "路由守卫负责登录状态和角色入口控制，最终授权仍由后端完成。",
        "查询组件在卸载或条件变化时取消失效请求，避免重复更新和无效网络调用。",
        "WebSocket 客户端只建立必要连接，并在退出登录或组件销毁时解除监听。",
        "明暗主题、响应式布局、低性能设备动效降级和减少动态效果偏好由公共层统一处理。",
    ])
    heading(doc, "3.6 后端模块设计", 2)
    add_body(doc, "后端采用 NestJS 模块化单体。控制器负责协议和输入边界，领域服务负责业务规则与事务，Prisma 负责持久化。跨域副作用通过明确服务调用触发，审计、通知和文件存储作为共享基础模块接入。")
    num, title, desc = FIGURE_PLACEHOLDERS[2]
    add_figure_placeholder(doc, num, title, desc)
    heading(doc, "3.7 数据与存储架构", 2)
    add_body(doc, "PostgreSQL 保存业务主数据、状态、关系和审计记录；Redis 保存需要跨实例共享的限流或短期状态；S3 兼容对象存储保存附件和资源。数据库只保存对象键、元数据和访问关系，文件下载必须经过授权接口。Prisma 迁移作为数据库结构版本基线。")
    heading(doc, "3.8 安全架构", 2)
    add_body(doc, "系统采用纵深防御：传输层使用 TLS；接入层配置安全响应头和受控 CORS；应用层使用 JWT、刷新令牌、RBAC、资源归属校验、输入验证及限流；数据层对凭据密钥加密并限制附件访问；治理层通过请求编号和审计日志追踪重要操作。访问令牌密钥与刷新令牌密钥均不得使用示例值，长度至少 32 字节且互不相同。")
    num, title, desc = FIGURE_PLACEHOLDERS[3]
    add_figure_placeholder(doc, num, title, desc)
    heading(doc, "3.9 外部依赖与降级", 2)
    add_table(doc, "3-3", "外部依赖与降级策略", ["依赖", "主要用途", "不可用时处理"], [
        ("SMTP 邮件服务", "重要通知、验证码和账户安全邮件", "站内通知与核心业务继续；失败写入审计，不自动重试"),
        ("AI 模型供应商", "报告建议、审核建议、知识增强问答", "返回统一错误；不改变漏洞状态，不阻断人工流程"),
        ("S3 兼容对象存储", "漏洞附件和学习资源", "拒绝依赖文件的提交或下载，避免产生失配元数据"),
        ("ClamAV", "上传文件恶意内容检测", "按环境安全策略拒绝或标记，具体生产降级策略未确认"),
        ("Redis", "共享限流和短期状态", "生产环境视为必需依赖，异常触发健康检查和运维处置"),
    ], [3.2, 5.2, 7.6])
    heading(doc, "3.10 可观测性与可靠性", 2)
    add_body(doc, "HTTP 请求通过请求编号关联前端错误提示和后端日志；高风险管理动作、登录安全事件、审核变更、邮件失败和 AI 调用均保留可审计记录。健康检查覆盖应用及关键依赖。关键数据库写入使用事务，避免报告状态、事件、积分或兑换发生部分成功。")


def add_chapter_4(doc):
    heading(doc, "4 功能需求业务流程设计说明", 1)
    heading(doc, "4.1 通用设计约束", 2)
    add_body(doc, "各功能模块统一遵循 DTO 输入校验、服务端权限判断、资源归属校验、结构化错误、请求编号、事务边界和审计要求。列表接口采用分页和稳定排序，文件接口不直接暴露对象存储内部地址。")
    for index, module in enumerate(MODULE_DESIGNS, 2):
        heading(doc, f"4.{index} {module['title']}", 2)
        add_table(doc, f"4-{index-1}", f"{module['title']}模块设计", ["设计项", "说明"], [
            ("核心组件", module["components"]),
            ("主要流程", module["flow"]),
            ("状态与持久化", module["state"]),
            ("权限边界", module["permission"]),
            ("异常与降级", module["failure"]),
        ], [3.4, 12.6])
        add_figure_placeholder(doc, f"4-{index-1}", f"{module['title']}设计图", module["figure"])
    heading(doc, f"4.{len(MODULE_DESIGNS) + 2} 接口与错误约定", 2)
    add_table(doc, f"4-{len(MODULE_DESIGNS) + 1}", "接口设计约定", ["项目", "约定"], [
        ("协议", "公开和业务接口使用 HTTPS REST；即时消息使用 Socket.IO WebSocket，并保留必要的 REST 历史查询。"),
        ("身份", "访问令牌通过统一认证机制传递；刷新令牌支持吊销和轮换。"),
        ("输入", "控制器 DTO 执行类型、长度、枚举、格式和业务前置条件校验。"),
        ("输出", "仅返回当前角色和资源所需字段；公开证书验证执行最小披露。"),
        ("错误", "返回稳定业务错误和请求编号；前端通过全局提示服务展示，不直接暴露内部堆栈。"),
        ("幂等与并发", "兑换、审核迁移、证书签发等关键操作通过唯一约束、状态条件和事务控制重复提交。"),
        ("审计", "高风险动作记录主体、动作、对象、结果、来源地址及必要上下文。"),
    ], [3.4, 12.6])


def add_chapter_5(doc, models):
    heading(doc, "5 数据库设计说明", 1)
    heading(doc, "5.1 设计原则", 2)
    add_body(doc, "数据库采用 PostgreSQL，Prisma Schema 是结构设计与迁移的事实基线。业务主表使用稳定主键和时间字段；关系、唯一性和状态约束尽量在数据库与服务端双重保证。附件正文与二进制文件分离，敏感凭据仅保存密文。")
    heading(doc, "5.2 数据域与模型清单", 2)
    rows = []
    for group, names in MODEL_GROUPS.items():
        rows.append((group, str(len(names)), "、".join(names)))
    add_table(doc, "5-1", "Prisma 数据模型分域清单", ["数据域", "模型数", "模型名称"], rows, [3.5, 1.8, 10.7])
    actual = set(models)
    declared = {name for names in MODEL_GROUPS.values() for name in names}
    if actual != declared:
        missing = sorted(actual - declared)
        stale = sorted(declared - actual)
        raise RuntimeError(f"model inventory mismatch; missing={missing}, stale={stale}")
    add_body(doc, f"当前 Prisma Schema 共定义 {len(actual)} 个模型。上述清单完整覆盖当前模型；字段级物理设计对核心业务模型展开，辅助日志、版本和关联模型以 Schema 及迁移文件为实现依据。")
    num, title, desc = FIGURE_PLACEHOLDERS[4]
    add_figure_placeholder(doc, num, title, desc)
    heading(doc, "5.3 核心物理数据模型", 2)
    add_body(doc, "字段类型按 Prisma 语法记录。带问号的字段允许为空，带方括号的字段为关系集合；默认值、唯一性、主键和关系约束以模型声明为准。表内英文模型名保持与代码一致，避免设计文档与实现映射失真。")
    for index, model_name in enumerate(CORE_MODELS, 2):
        model = models[model_name]
        heading(doc, f"5.3.{index-1} {model_name}", 3)
        rows = []
        for seq, (field, field_type, attrs) in enumerate(model["fields"], 1):
            rows.append((seq, field, field_type, constraint_text(field_type, attrs), field_description(field, field_type)))
        add_table(doc, f"5-{index}", f"{model_name} 字段设计", ["序号", "字段", "Prisma 类型", "约束", "说明"], rows, [1.2, 3.2, 2.6, 4.0, 5.0])
        if model["indexes"]:
            add_body(doc, "模型级约束与索引：" + "；".join(model["indexes"]) + "。")
    heading(doc, "5.4 一致性、索引与事务", 2)
    add_bullets(doc, [
        "登录标识、证书编号、权限编码、对象存储键和需要去重的业务标识使用唯一约束。",
        "面向状态、用户、创建时间和更新时间的高频列表查询建立组合索引，具体定义以 Prisma Schema 为准。",
        "审核状态迁移、积分流水、商城兑换和证书签发在事务内完成，并以当前状态作为并发更新条件。",
        "通知、AI 调用和审计等副作用根据业务重要性处理；不得因非关键外部服务失败破坏核心状态。",
    ])
    heading(doc, "5.5 迁移、备份与保留", 2)
    add_body(doc, "数据库变更通过 Prisma 迁移版本化，部署前执行迁移检查，禁止在正式环境手工修改结构而不形成迁移记录。备份应同时覆盖 PostgreSQL 与对象存储，并验证可恢复性。聊天容量达到数据库容量阈值时提示管理员按日期清理并可先导出备份；正式备份周期、保留期限和恢复目标未确认，应由部署运维方案另行确定。")
    heading(doc, "5.6 数据安全", 2)
    add_body(doc, "密码只保存摘要；模型 API Key 使用独立的 AI_CREDENTIAL_ENCRYPTION_KEY 加密；访问令牌与刷新令牌使用不同密钥。对象存储文件通过业务授权访问，AI 分析仅接收必要的脱敏内容。审计和日志不得记录完整令牌、密码、模型密钥或未经处理的敏感漏洞数据。")


def add_chapter_6(doc):
    heading(doc, "6 软件界面设计", 1)
    heading(doc, "6.1 界面设计原则", 2)
    add_body(doc, "界面以数据和任务为中心，保持克制、清晰和一致。公共图标只承担明确操作或状态含义，不作为装饰。所有页面复用统一的色彩、字体、间距、表单、弹窗、加载动画和全局提示服务，并支持明暗主题。")
    add_table(doc, "6-1", "界面公共规范", ["项目", "设计要求"], [
        ("布局", "内容宽度以管理员首页为统一基准；工作台采用可扫描的信息密度和稳定网格。"),
        ("反馈", "成功、错误、警告和处理中状态统一使用全局提示服务；重要操作使用二次确认。"),
        ("表单", "输入框提供长度、格式、必填和枚举约束；服务端继续执行同等或更严格校验。"),
        ("弹窗", "详情、模型配置、Skill、附件和 AI 建议使用统一弹窗层级并覆盖导航交互。"),
        ("响应式", "侧栏在可能遮挡主内容前自动折叠；固定格式组件具有稳定尺寸约束。"),
        ("主题", "所有文本、边框、输入、卡片和遮罩均适配明暗主题，保证对比度。"),
        ("动效", "动效用于状态过渡和空间关系；低性能设备和减少动态效果偏好下降级。"),
        ("加载", "可见内容和关键字体未就绪时显示全局圆形加载，完成后淡出。"),
    ], [3.2, 12.8])
    heading(doc, "6.2 页面区域设计", 2)
    add_table(doc, "6-2", "页面区域与设计重点", ["页面区域", "设计重点"], UI_AREAS, [4.0, 12.0])
    for index, (area, description) in enumerate(UI_AREAS, 1):
        heading(doc, f"6.2.{index} {area}", 3)
        add_body(doc, description)
        add_figure_placeholder(doc, f"6-{index}", f"{area}界面示意图", f"界面截图占位。后续使用当前测试环境的真实页面截图，保持原始宽高比，不随意裁剪或压缩；截图应覆盖“{area}”的核心任务、主要状态和必要操作，并遮蔽账号、漏洞详情、密钥及其他敏感信息。")
    heading(doc, "6.3 响应式与可访问性", 2)
    add_body(doc, "界面在桌面、平板和移动视口下保持主要任务可达。文本不得与控件重叠，按钮和输入框具有可用触控区域；键盘焦点、表单标签、对话框语义和错误提示应可被辅助技术识别。颜色不作为唯一状态表达方式。")
    heading(doc, "6.4 性能与资源策略", 2)
    add_body(doc, "首屏优先加载必要路由、字体和可见资源；非首屏模块和 Three.js 场景按需加载。动画避免触发布局抖动，WebGL 对象在卸载时释放；不可见页面停止无效渲染和轮询。图片保持视觉质量并使用适当格式、尺寸和缓存策略。")
    heading(doc, "6.5 界面验收要点", 2)
    add_bullets(doc, [
        "公开、用户、审核员和管理员路由在授权边界内可达，越权入口不可操作。",
        "明暗主题下文字、输入、表格、弹窗和 AI 可操作卡片均清晰可读。",
        "全局提示、加载、确认和错误恢复方式一致，错误包含可追踪请求编号。",
        "移动与窄屏下侧栏、表格、弹窗和固定操作区不遮挡核心内容。",
        "真实数据、空状态、加载状态、错误状态和权限不足状态均有明确呈现。",
    ])
    heading(doc, "6.6 待补工程图与截图清单", 2)
    rows = []
    for number, title, description in FIGURE_PLACEHOLDERS:
        rows.append((f"图 {number}", title, description))
    for index, (area, _) in enumerate(UI_AREAS, 1):
        rows.append((f"图 6-{index}", f"{area}界面示意图", "使用测试环境真实截图，保持比例并完成敏感信息遮蔽。"))
    for index, module in enumerate(MODULE_DESIGNS, 1):
        rows.append((f"图 4-{index}", f"{module['title']}设计图", module["figure"]))
    add_table(doc, "6-3", "待补工程图与截图清单", ["图号", "图名", "绘制或取图要求"], rows, [2.2, 4.4, 9.4])


def build_document():
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    doc = Document(REFERENCE)
    remove_body(doc)
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        clear_header_footer(section)
    add_cover(doc)
    add_control_pages(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(body_section)
    add_header_footer(body_section)
    add_toc(doc)
    page_break(doc)
    models = parse_prisma_models()
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc, models)
    add_chapter_6(doc)
    normalize_table_paragraphs(doc)
    doc.core_properties.title = f"{PROJECT}项目系统设计说明书"
    doc.core_properties.subject = "软件系统设计"
    doc.core_properties.author = "蓝山工作室运维安全部安全组"
    doc.core_properties.keywords = "校园漏洞响应, 安全学习, 系统设计, 软件工程"
    doc.core_properties.comments = "依据当前仓库实现生成"
    set_update_fields_on_open(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    deduplicate_package(OUTPUT)
    return doc, models


if __name__ == "__main__":
    document, parsed_models = build_document()
    print(f"generated: {OUTPUT}")
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} models={len(parsed_models)}")
