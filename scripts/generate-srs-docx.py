#!/usr/bin/env python3
"""Generate the formal SRS from the preserved project Word template."""

from __future__ import annotations

from copy import deepcopy
from datetime import date
import os
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.image.image import Image as DocxImage
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from assets.srs_diagram_specs import USE_CASE_DIAGRAM_SPECS


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "校园漏洞响应与安全学习平台软件需求规格说明书.docx"
USE_CASE_REFERENCE = ROOT / "scripts" / "assets" / "srs-use-case-reference.docx"
OUTPUT = TEMPLATE
DIAGRAM_DIR = ROOT / "docs" / "assets" / "diagrams"
FIGURE_3_1 = DIAGRAM_DIR / "figure-3-1-vulnerability-response.png"
FIGURE_3_2 = DIAGRAM_DIR / "figure-3-2-ai-knowledge-flow.png"
FIGURE_4_1 = DIAGRAM_DIR / "figure-4-1-system-use-cases.png"
FIGURE_4_2_ACTIVITY = DIAGRAM_DIR / "figure-4-2a-public-certificate-activity.png"
FIGURE_4_2_SEQUENCE = DIAGRAM_DIR / "figure-4-2b-public-certificate-sequence.png"
GENERATED_ON = date(2026, 8, 26)
PROJECT_NAME = "校园漏洞响应与安全学习平台"
VERSION = "V2.10"
_REFERENCE_DOCUMENT = None


def reference_document() -> Document:
    global _REFERENCE_DOCUMENT
    if _REFERENCE_DOCUMENT is None:
        _REFERENCE_DOCUMENT = Document(USE_CASE_REFERENCE)
    return _REFERENCE_DOCUMENT


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="000000", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_borders(cell, color="000000", size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_fixed_table_layout(table, width_twips: int) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.insert(0, tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_run_font(run, name="宋体", size=12, bold=None, color=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction: str, fallback: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text_node, end])


def remove_body_content(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def style_named(doc: Document, name: str):
    for style in doc.styles:
        if style.name == name:
            return style
    raise KeyError(f"no style with name {name!r}")


def has_style(doc: Document, name: str) -> bool:
    return any(style.name == name for style in doc.styles)


def configure_styles(doc: Document) -> None:
    normal = style_named(doc, "Normal")
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    body_style = style_named(doc, "正文缩进") if has_style(doc, "正文缩进") else normal
    body_style.font.name = "宋体"
    body_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    body_style.font.size = Pt(12)
    body_style.paragraph_format.first_line_indent = Pt(24)
    body_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body_style.paragraph_format.space_after = Pt(0)
    body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = style_named(doc, style_name)
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)
    page_number = section._sectPr.find(qn("w:pgNumType"))
    if page_number is not None:
        section._sectPr.remove(page_number)


def clear_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for container in (section.header, section.footer):
        for paragraph in list(container.paragraphs):
            paragraph._element.getparent().remove(paragraph._element)


def configure_running_header_footer(section) -> None:
    clear_header_footer(section)

    header = section.header
    hp = header.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hp.add_run(f"{PROJECT_NAME}软件需求规格说明书    版本：{VERSION}"), "宋体", 9)
    hp.paragraph_format.space_after = Pt(0)

    footer = section.footer
    fp = footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("第 "), "宋体", 9)
    add_field(fp, "PAGE", "1")
    set_run_font(fp.add_run(" 页，共 "), "宋体", 9)
    add_field(fp, "NUMPAGES", "1")
    set_run_font(fp.add_run(" 页"), "宋体", 9)


def configure_section(doc: Document) -> None:
    section = doc.sections[0]
    configure_page(section)
    clear_header_footer(section)


def add_body(doc: Document, text: str, *, indent=True, bold_prefix: str | None = None):
    style = style_named(doc, "正文缩进") if indent and has_style(doc, "正文缩进") else style_named(doc, "Normal")
    p = doc.add_paragraph(style=style)
    if not indent:
        p.paragraph_format.first_line_indent = None
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        set_run_font(p.add_run(f"{idx}. {item}"))


def add_heading(doc: Document, text: str, level: int):
    if level <= 2:
        text = re.sub(r"^\d+(?:\.\d+)*\s*", "", text)
    p = doc.add_paragraph(style=style_named(doc, f"Heading {level}"))
    if level == 1:
        p.paragraph_format.page_break_before = True
    set_run_font(p.add_run(text), "黑体", {1: 16, 2: 14, 3: 12}.get(level, 12), True)
    keep_with_next(p)
    return p


def add_control_heading(doc: Document, text: str):
    p = doc.add_paragraph(style=style_named(doc, "Normal"))
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), "黑体", 14, True)
    return p


def clone_reference_table(doc: Document, reference_index: int):
    source = reference_document().tables[reference_index]
    placeholder = doc.add_table(rows=1, cols=1)
    placeholder._tbl.getparent().replace(placeholder._tbl, deepcopy(source._tbl))
    return doc.tables[-1]


def remove_last_table_column(table) -> None:
    grid_columns = table._tbl.tblGrid.gridCol_lst
    if len(grid_columns) < 2:
        raise ValueError("Cannot remove the last table column")
    removed_width = int(grid_columns[-1].get(qn("w:w")) or 0)
    previous_width = int(grid_columns[-2].get(qn("w:w")) or 0)
    grid_columns[-2].set(qn("w:w"), str(previous_width + removed_width))
    table._tbl.tblGrid.remove(grid_columns[-1])
    for row in table._tbl.tr_lst:
        cells = row.tc_lst
        previous_cell, removed_cell = cells[-2], cells[-1]
        previous_tc_w = previous_cell.tcPr.find(qn("w:tcW"))
        removed_tc_w = removed_cell.tcPr.find(qn("w:tcW"))
        if previous_tc_w is not None and removed_tc_w is not None:
            width = int(previous_tc_w.get(qn("w:w")) or 0) + int(removed_tc_w.get(qn("w:w")) or 0)
            previous_tc_w.set(qn("w:w"), str(width))
        row.remove(removed_cell)


def resize_reference_table(table, row_count: int, column_count: int) -> None:
    while len(table.columns) > column_count:
        remove_last_table_column(table)
    if len(table.columns) != column_count:
        raise ValueError(f"Reference table cannot provide {column_count} columns")
    body_template = deepcopy(table.rows[1]._tr if len(table.rows) > 1 else table.rows[0]._tr)
    while len(table.rows) < row_count:
        table._tbl.append(deepcopy(body_template))
    while len(table.rows) > row_count:
        table._tbl.remove(table.rows[-1]._tr)


def set_table_column_widths(table, widths_cm: list[float]) -> None:
    if len(widths_cm) != len(table.columns):
        raise ValueError("Table width count does not match column count")
    widths_twips = [str(Cm(width).twips) for width in widths_cm]
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths_twips):
        grid_column.set(qn("w:w"), width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), width)
            cell_width.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def allow_automatic_row_height(row) -> None:
    """Remove template row-height constraints so borders follow cell content."""
    row_properties = row._tr.get_or_add_trPr()
    for height in list(row_properties.findall(qn("w:trHeight"))):
        row_properties.remove(height)


def add_table_caption(doc: Document, number: str, title: str) -> None:
    caption = doc.add_paragraph(style=style_named(doc, "Normal"))
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True
    set_run_font(caption.add_run(f"表 {number}  {title}"), "宋体", 10.5)


def add_table(
    doc: Document,
    number: str,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    widths=None,
    reference_index=None,
):
    add_table_caption(doc, number, title)
    column_count = len(headers)
    width_presets = {
        ("编号", "资料名称", "版本或日期", "发布单位", "用途"): [1.2, 5.0, 2.8, 3.2, 3.6],
        ("术语", "英文或全称", "说明"): [1.8, 4.4, 9.6],
        ("编号", "业务需求", "验收口径", "优先级"): [1.5, 5.8, 7.0, 1.5],
        ("编号", "风险", "影响", "控制措施"): [1.5, 3.5, 4.7, 6.1],
        ("编号", "假设或依赖", "约束与应对"): [1.5, 6.1, 8.2],
        ("编号", "涉众", "主要职责或行为", "核心期望"): [1.4, 2.4, 6.0, 6.0],
        ("编号", "用例名称", "主要参与者", "优先级"): [1.4, 4.8, 6.4, 3.2],
        ("编号", "需求", "验收方法"): [2.2, 6.8, 6.8],
        ("优先级", "范围", "说明"): [2.0, 6.8, 7.0],
        ("图号", "图名", "主要用途", "主要内容"): [1.5, 4.2, 3.2, 6.9],
    }
    if widths is None:
        widths = width_presets.get(tuple(headers))
    if reference_index is None:
        reference_index = {2: 3, 3: 16, 4: 16, 5: 4}.get(column_count)
    if reference_index is None:
        raise ValueError(f"No reference table for {column_count} columns")
    table = clone_reference_table(doc, reference_index)
    resize_reference_table(table, 1 + len(rows), column_count)
    if widths is not None:
        set_table_column_widths(table, widths)
    set_repeat_table_header(table.rows[0])
    for column_index, header in enumerate(headers):
        cell = table.rows[0].cells[column_index]
        replace_cell_paragraphs(cell, [(cell.paragraphs[0], str(header), None)])
    for row_index, row_data in enumerate(rows, 1):
        prevent_row_split(table.rows[row_index])
        for column_index, value in enumerate(row_data):
            cell = table.rows[row_index].cells[column_index]
            replace_cell_paragraphs(cell, [(cell.paragraphs[0], str(value), None)])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_cover_table(doc: Document, values: list[str]):
    table = clone_reference_table(doc, 0)
    if len(values) != len(table.rows):
        raise ValueError("Cover table value count does not match reference")
    labels = ["项目名称", "产品类型", "项目编号", "版本号", "文档状态", "编制人", "所属部门", "提交日期"]
    for row, label, value in zip(table.rows, labels, values):
        replace_cell_paragraphs(row.cells[0], [(row.cells[0].paragraphs[0], label, None)])
        cell = row.cells[1]
        replace_cell_paragraphs(cell, [(cell.paragraphs[0], value, None)])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_key_value_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(3.4)
        cells[1].width = Cm(13.0)
        for cell in cells:
            set_cell_margins(cell, 0, 108, 0, 108)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cells[0], "E5E5E5")
        cells[0].text = ""
        cells[1].text = ""
        set_run_font(cells[0].paragraphs[0].add_run(label), "黑体", 10.5, True)
        set_run_font(cells[1].paragraphs[0].add_run(value), "宋体", 10.5)
        cells[0].paragraphs[0].paragraph_format.line_spacing = 1.2
        cells[1].paragraphs[0].paragraph_format.line_spacing = 1.2
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def replace_paragraph_text(template_paragraph, text: str, num_id: int | None = None):
    paragraph = deepcopy(template_paragraph._p)
    run_template = next((child for child in paragraph if child.tag == qn("w:r")), None)
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    run = deepcopy(run_template) if run_template is not None else OxmlElement("w:r")
    for child in list(run):
        if child.tag != qn("w:rPr"):
            run.remove(child)
    text_node = OxmlElement("w:t")
    text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    paragraph.append(run)
    if num_id is not None:
        p_pr = paragraph.find(qn("w:pPr"))
        num_pr = p_pr.find(qn("w:numPr")) if p_pr is not None else None
        num_id_node = num_pr.find(qn("w:numId")) if num_pr is not None else None
        if num_id_node is not None:
            num_id_node.set(qn("w:val"), str(num_id))
    return paragraph


def replace_cell_paragraphs(cell, entries: list[tuple[object, str, int | None]]) -> None:
    for paragraph in list(cell._tc.findall(qn("w:p"))):
        cell._tc.remove(paragraph)
    for template_paragraph, text, num_id in entries:
        cell._tc.append(replace_paragraph_text(template_paragraph, text, num_id))


def clone_reference_numbering(doc: Document, reference: Document, source_id=14) -> int:
    source_root = reference.part.numbering_part.element
    target_root = doc.part.numbering_part.element
    source_num = next(node for node in source_root.findall(qn("w:num")) if node.get(qn("w:numId")) == str(source_id))
    source_abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    source_abstract = next(
        node for node in source_root.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) == source_abstract_id
    )
    existing_abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in target_root.findall(qn("w:abstractNum"))]
    existing_num_ids = [int(node.get(qn("w:numId"))) for node in target_root.findall(qn("w:num"))]
    new_abstract_id = max(existing_abstract_ids, default=0) + 1
    new_num_id = max(existing_num_ids, default=0) + 1

    abstract_copy = deepcopy(source_abstract)
    abstract_copy.set(qn("w:abstractNumId"), str(new_abstract_id))
    first_num = target_root.find(qn("w:num"))
    target_root.insert(target_root.index(first_num) if first_num is not None else len(target_root), abstract_copy)

    num_copy = deepcopy(source_num)
    num_copy.set(qn("w:numId"), str(new_num_id))
    num_copy.find(qn("w:abstractNumId")).set(qn("w:val"), str(new_abstract_id))
    target_root.append(num_copy)
    return new_num_id


def add_use_case_table(doc: Document, number: str, use_case: dict, related: str) -> None:
    """Clone the reference use-case table XML and replace content only."""
    add_table_caption(doc, number, f"{use_case['name']}用例规约")
    reference = Document(USE_CASE_REFERENCE)
    reference_table = reference.tables[6]
    placeholder = doc.add_table(rows=1, cols=1)
    placeholder._tbl.getparent().replace(placeholder._tbl, deepcopy(reference_table._tbl))
    table = doc.tables[-1]
    list_num_id = clone_reference_numbering(doc, reference)

    summary_values = [
        f"{use_case['id']} {use_case['name']}",
        use_case["summary"],
        use_case["actors"],
        related,
        use_case["pre"],
        use_case["post"],
    ]
    for row_index, value in enumerate(summary_values):
        value_cell = table.rows[row_index].cells[1]
        template_paragraph = value_cell.paragraphs[0]
        replace_cell_paragraphs(value_cell, [(template_paragraph, value, None)])

    basic_cell = table.rows[6].cells[0]
    basic_templates = basic_cell.paragraphs
    replace_cell_paragraphs(
        basic_cell,
        [(basic_templates[0], "基本事件流", None)]
        + [(basic_templates[1], step, list_num_id) for step in use_case["main"]],
    )

    alternate_cell = table.rows[7].cells[0]
    alternate_templates = alternate_cell.paragraphs
    replace_cell_paragraphs(
        alternate_cell,
        [(alternate_templates[0], "备选事件流", None)]
        + [(alternate_templates[1], f"A-{index} {step}", None) for index, step in enumerate(use_case["alt"], 1)],
    )

    constraint_cell = table.rows[8].cells[0]
    constraint_templates = constraint_cell.paragraphs
    constraint_entries = [
        (constraint_templates[0], "补充约束-数据需求", None),
        (constraint_templates[1], f"D-1 {use_case['data']}", None),
        (constraint_templates[3], "补充约束-业务规则", None),
    ]
    for index, rule in enumerate(use_case["rules"], 1):
        normalized_rule = re.sub(r"^BR-[0-9-]+\s*", "", rule)
        constraint_entries.append((constraint_templates[4], f"B-{index} {normalized_rule}", None))
    replace_cell_paragraphs(constraint_cell, constraint_entries)

    issue_cell = table.rows[9].cells[0]
    issue_templates = issue_cell.paragraphs
    replace_cell_paragraphs(issue_cell, [
        (issue_templates[0], "待解决问题", None),
        (issue_templates[1], "无", None),
    ])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure_placeholder(doc: Document, number: str, title: str, description: str) -> None:
    figures = {
        "3-1": [(FIGURE_3_1, f"图 {number}  {title}")],
        "3-2": [(FIGURE_3_2, f"图 {number}  {title}")],
        "4-1": [(FIGURE_4_1, f"图 {number}  {title}")],
        "4-2": [
            (FIGURE_4_2_ACTIVITY, "图 4-2(a)  公开访问与证书验证活动图"),
            (FIGURE_4_2_SEQUENCE, "图 4-2(b)  公开访问与证书验证系统顺序图"),
        ],
    }
    for spec in USE_CASE_DIAGRAM_SPECS:
        spec_number = spec["number"]
        slug = spec["slug"]
        spec_title = spec["title"]
        figures[spec_number] = [
            (DIAGRAM_DIR / f"figure-{spec_number}-{slug}-activity.png", f"图 {spec_number}(a)  {spec_title}活动图"),
            (DIAGRAM_DIR / f"figure-{spec_number}-{slug}-sequence.png", f"图 {spec_number}(b)  {spec_title}系统顺序图"),
        ]
    available_figures = figures.get(number)
    if available_figures and all(path.exists() for path, _ in available_figures):
        for path, caption in available_figures:
            figure = doc.add_paragraph()
            figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
            figure.paragraph_format.page_break_before = True
            figure.paragraph_format.keep_with_next = True
            figure.paragraph_format.space_before = Pt(3)
            figure.paragraph_format.space_after = Pt(0)
            image = DocxImage.from_file(str(path))
            aspect_ratio = image.px_width / image.px_height
            max_width = 15.5 if number == "4-1" else 14.2
            max_height = 22.0
            figure_width = min(max_width, max_height * aspect_ratio)
            figure.add_run().add_picture(str(path), width=Cm(figure_width))
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = False
            cap.paragraph_format.space_after = Pt(6)
            set_run_font(cap.add_run(caption), "宋体", 10.5)
        return

    table = clone_reference_table(doc, 3)
    resize_reference_table(table, 1, 2)
    cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    set_table_borders(table, color="C0C0C0", size="4")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, 180, 180, 180, 180)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(f"【图示占位：{title}】\n"), "黑体", 11, True)
    set_run_font(p.add_run(description), "宋体", 10.5)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    set_run_font(cap.add_run(f"图 {number}  {title}（待绘制）"), "宋体", 10.5)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(p.add_run("受控文件"), "宋体", 10)
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(PROJECT_NAME), "黑体", 24, True)
    p.paragraph_format.space_after = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("软件需求规格说明书"), "黑体", 22, True)
    for _ in range(4):
        doc.add_paragraph()
    add_cover_table(doc, [
        PROJECT_NAME,
        "Web 应用系统",
        "",
        VERSION,
        "受控稿",
        "",
        "",
        GENERATED_ON.strftime("%Y年%m月%d日"),
    ])
    add_page_break(doc)


def build_document_control(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("文档控制"), "黑体", 18, True)
    add_control_heading(doc, "文档更新记录")
    add_table(doc, "0-1", "文档更新记录", ["日期", "更新人", "版本", "备注"], [[GENERATED_ON.isoformat(), "", VERSION, "依据现行业务边界全面修订功能、用例、数据约束及非功能需求"]], reference_index=1)
    add_control_heading(doc, "文档审核记录")
    add_table(doc, "0-2", "文档审核记录", ["日期", "审核人", "职务", "审核版本", "备注"], [["", "", "", VERSION, ""]], reference_index=2)
    add_control_heading(doc, "文档发行范围")
    add_table(doc, "0-3", "文档发行范围", ["接收部门或人员", "用途"], [
        ["项目建设与管理单位", "项目立项、范围控制、交付验收和变更管理"],
        ["系统设计、开发与测试人员", "系统设计、开发实现、测试验证和缺陷判定"],
        ["平台运营与安全管理人员", "制度配置、业务运行、安全管理和审计监督"],
    ], reference_index=3)


def build_toc(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(section)
    configure_running_header_footer(section)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目  录"), "黑体", 18, True)
    toc = doc.add_paragraph()
    toc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(toc.add_run("[[SRS_TOC]]"), "宋体", 12)


def build_intro(doc: Document) -> None:
    add_heading(doc, "1 引言", 1)
    add_heading(doc, "1.1 编写目的", 2)
    add_body(doc, "本文档规定校园漏洞响应与安全学习平台的业务目标、系统边界、参与者职责、功能需求、业务规则、数据约束及非功能要求，建立统一、可验证、可追踪的软件需求基线。本文档作为系统架构设计、详细设计、开发实现、测试验证、部署交付、运行维护和需求变更控制的依据。")

    add_heading(doc, "1.2 预期读者和阅读建议", 2)
    add_body(doc, "本文档面向建设单位、项目管理人员、产品与架构设计人员、软件研发人员、测试与质量保证人员、部署运维人员、平台管理员、审核员及安全管理责任人员。建设与管理人员宜重点阅读第2章和第3章，以掌握业务目标、建设范围、总体方案和业务风险；设计、研发与测试人员宜重点阅读第4章和第5章，以建立功能、状态、权限、数据约束及验收基线；部署运维人员还应关注可靠性、安全性、备份恢复、容量和外部依赖要求。所有参与方应统一采用本文档定义的术语、需求编号、优先级和业务规则。")

    add_heading(doc, "1.3 文档约定", 2)
    add_bullets(doc, [
        "正文采用宋体小四、1.5倍行距；各级标题、表格和图题按模板样式编排。",
        "“应”表示必须满足的需求；“宜”表示推荐满足的需求；“可”表示在条件允许时提供的能力。",
        "需求使用 FR（功能需求）、BR（业务规则）和 NFR（非功能需求）编号；P0、P1、P2依次表示核心、重要和增强优先级。优先级用于规定业务重要性与交付顺序。",
        "本文所称“AI”均为辅助能力，不具备代替人工审核、风险定级、证书签发或其他管理决定的权限。",
        "本文图示采用可编辑源文件生成，并同步提供SVG与PNG格式。活动图、泳道图、用例图和系统顺序图应与对应需求规约保持一致；需求变更涉及参与者、流程、状态或外部依赖时，应同步更新图稿、题注和图示索引。",
        "业务成效指标应以平台审计数据、运行报表或经确认的统计资料为依据，不采用无法追溯的数据。",
    ])

    add_heading(doc, "1.4 参考资料", 2)
    add_table(doc, "1-1", "参考资料", ["编号", "资料名称", "版本或日期", "发布单位", "用途"], [
        ["R-01", "《中华人民共和国网络安全法》", "2025年修正，2026年1月1日起施行", "全国人民代表大会常务委员会", "网络运行安全与安全保护义务依据"],
        ["R-02", "《中华人民共和国数据安全法》", "2021年", "全国人民代表大会常务委员会", "数据处理与安全保护依据"],
        ["R-03", "《中华人民共和国个人信息保护法》", "2021年", "全国人民代表大会常务委员会", "个人信息最小必要处理依据"],
        ["R-04", "GB/T 8567—2006《计算机软件文档编制规范》", "2006年", "国家标准化管理委员会", "软件文档结构参考"],
        ["R-05", "GB/T 9385—2008《计算机软件需求规格说明规范》", "2008年", "国家标准化管理委员会", "需求内容、特性与编制规则依据"],
        ["R-06", "《系统总体设计说明书》", VERSION, "项目建设单位", "系统边界、总体架构和模块关系依据"],
        ["R-07", "《数据库设计说明书》", VERSION, "项目建设单位", "数据实体、状态、关系和约束依据"],
        ["R-08", "《接口设计说明书》", VERSION, "项目建设单位", "系统接口、数据交换和异常处理依据"],
    ])
    add_body(doc, "法律法规与国家标准的适用版本以有权机关及国家标准信息公共服务平台公布的信息为准。本文档仅提炼与系统需求相关的约束，不替代专项法律意见或合规审查结论。", indent=False)

    add_heading(doc, "1.5 缩写术语", 2)
    add_table(doc, "1-2", "缩写术语", ["术语", "英文或全称", "说明"], [
        ["SRC", "Security Response Center", "接收、审核、处置、跟踪和反馈安全漏洞的响应机制。"],
        ["RBAC", "Role-Based Access Control", "基于角色和权限点控制系统访问。"],
        ["AI服务网关", "Artificial Intelligence Service Gateway", "统一承接模型调用、上下文组织、去标识化、审计与降级处理的可信服务边界。"],
        ["Skill", "声明式技能", "向模型提供特定任务说明、边界和输出约束的可复用配置，不直接授予系统权限。"],
        ["知识库", "AI Knowledge Base", "由经审核或规则准入的可复用问答知识构成，用于相似问题检索与回答增强，不转换为Skill。"],
        ["RPO", "Recovery Point Objective", "系统故障后可接受的数据丢失时间上限。"],
        ["RTO", "Recovery Time Objective", "系统故障后恢复业务所允许的最长时间。"],
        ["P95", "95th Percentile", "95%的请求响应时间不超过给定阈值。"],
        ["去标识化", "De-identification", "在不借助额外信息时降低对特定自然人或资产识别能力的处理。"],
    ])


def build_business(doc: Document) -> None:
    add_heading(doc, "2 业务需求", 1)
    add_heading(doc, "2.1 项目背景及应用场景", 2)
    add_body(doc, "校园信息系统数量持续增加，漏洞发现者、审核人员、系统责任方与平台管理人员之间需要在合法授权边界内完成报告提交、证据补充、审核确认、修复协作、复测归档和结果反馈。若主要依赖即时通信、邮件或分散表格，容易出现材料格式不统一、处理状态不可追踪、责任边界不清、敏感信息扩散、贡献记录难以核验等问题。与此同时，初学者需要在参与安全测试前理解网络安全法律法规、授权范围与禁止行为，并在实践后获得可验证的学习资源、积分记录和证书反馈。")
    add_body(doc, "本平台面向单一校园部署场景，服务对象包括未登录访客、完成注册的普通用户、承担审核与复测职责的审核员，以及负责平台配置、权限、证书、内容和审计的管理员。平台以漏洞响应全生命周期为主线，将安全准入、Markdown漏洞报告、附件证据、审核处置、通知提醒、证书验证、积分激励、学习资源和受控AI辅助纳入统一系统。典型场景包括：用户在授权资产范围内发现问题后提交报告；审核员要求补充材料或给出人工结论；管理员配置责任书、题库、资产、通知、证书和AI策略；访客通过证书编号核验证书状态。")

    add_heading(doc, "2.2 业务机遇", 2)
    add_body(doc, "项目的核心价值在于将分散的漏洞处置活动转化为可追踪、可审计、可反馈的校园安全协作流程。与通用工单系统相比，本平台将安全准入、漏洞专业字段、证据附件、审核处置、复测归档和证书激励统一建模；与单一安全学习系统相比，本平台通过授权范围内的实践活动形成学习、提交、审核和反馈闭环；与完全依赖外部AI的方案相比，本平台通过可信服务边界、数据去标识化、人工决策边界、知识检索和按需调用控制安全风险与运行成本。平台由此形成兼顾漏洞响应、安全教育、贡献激励和治理审计的综合业务能力。")

    add_heading(doc, "2.3 业务需求", 2)
    add_table(doc, "2-1", "业务需求", ["编号", "业务需求", "验收口径", "优先级"], [
        ["BRQ-01", "建立漏洞提交、审核、补充、确认、修复、复测、归档的闭环。", "每份正式报告具有唯一编号、当前状态和完整事件记录。", "P0"],
        ["BRQ-02", "在提交前落实安全责任、知识答题和授权边界。", "未满足准入条件的普通用户不能正式提交漏洞。", "P0"],
        ["BRQ-03", "对关键状态变化提供站内通知和必要的邮件提醒。", "要求补充、驳回、确认有效、归档和证书签发均形成可追踪通知。", "P0"],
        ["BRQ-04", "建立证书签发、撤销和公开验证机制。", "证书编号唯一；公开查询只返回核验所需信息和当前状态。", "P0"],
        ["BRQ-05", "形成可核验的积分账本、排行榜和兑换流程。", "积分余额与账本一致；变更原因、时间和关联业务可追溯。", "P1"],
        ["BRQ-06", "提供学习资源、公告和受控沟通能力。", "用户可获得与平台运行相关的内容，管理员可维护并审计。", "P1"],
        ["BRQ-07", "在不替代人工决策的前提下提供AI辅助。", "AI调用按需触发，结果标注辅助性质，关键决定仍由授权人员完成。", "P2"],
        ["BRQ-08", "保证角色权限、敏感数据访问和关键操作可审计。", "高风险管理与审核操作审计覆盖率目标为100%。", "P0"],
    ])

    add_heading(doc, "2.4 业务风险", 2)
    add_table(doc, "2-2", "业务风险", ["编号", "风险", "影响", "控制措施"], [
        ["RISK-01", "未授权测试或越界操作", "影响校园系统稳定，带来合规和安全责任。", "强制准入、责任书、授权资产范围、操作提示和审计追踪。"],
        ["RISK-02", "漏洞与个人敏感信息泄露", "扩大攻击面并损害用户权益。", "最小权限、受控附件、去标识化、日志脱敏、必要信息邮件。"],
        ["RISK-03", "AI输出错误或被过度信任", "误导报告完善或审核判断。", "明确辅助边界、来源与限制提示、人工复核、点踩转人工检查。"],
        ["RISK-04", "外部模型、邮箱或对象存储不可用", "部分增强功能中断。", "隔离依赖、超时与降级、失败审计，不阻塞无关核心流程。"],
        ["RISK-05", "需求与管理制度变更未同步", "系统规则与实际管理要求不一致。", "建立需求变更、配置复核、版本发布和制度映射机制。"],
        ["RISK-06", "权限配置错误或高权限账号滥用", "导致越权查看、修改或泄露。", "RBAC最小权限、高风险确认、异地登录提醒与完整审计。"],
    ])


def build_vision(doc: Document) -> None:
    add_heading(doc, "3 项目前景", 1)
    add_heading(doc, "3.1 前景概述", 2)
    add_body(doc, "平台拟形成适用于单一校园环境的漏洞响应与安全学习基础设施。其长期目标是让授权范围内的安全发现能够进入统一、规范且可追踪的处置路径，让普通用户清楚知道是否需要补充材料、审核进行到何处、获得了何种积分或证书以及下一步应做什么；让审核员在清晰权限边界内完成复现、定级建议、补充要求和复测归档；让管理员能够配置制度、题库、资产、证书、通知、AI与知识库策略，并通过审计日志发现异常。平台不追求用自动化取代专业判断，而是通过结构化数据、状态机、受控沟通和辅助分析降低协作成本。未来演进应坚持核心业务不依赖单一外部模型或供应商，并保证历史记录、权限边界和审计证据在功能扩展中保持稳定。")

    add_heading(doc, "3.2 主要解决方案", 2)
    add_body(doc, "经比较，通用工单系统难以完整表达漏洞报告、复测、证书和准入规则；仅建设漏洞数据库无法提供流程协作与反馈；完全依赖外部AI又会引入成本、隐私和可用性风险。因此，项目采用“结构化业务平台为核心、AI按需辅助、知识库检索增强、外部能力可替换”的组合方案。")

    add_heading(doc, "3.2.1 以状态机和权限控制构建漏洞响应闭环", 3)
    add_body(doc, "系统以漏洞报告为核心聚合基本信息、Markdown正文、原始附件、补充材料、审核记录、复测记录、状态事件、积分和证书。服务端应依据角色、权限点、资源归属及状态前置条件校验每次操作；关键变更应在同一事务边界内同步更新业务状态、事件和通知。客户端应依据角色呈现授权范围内的业务入口，减少重复录入，但不得以界面隐藏替代服务端授权控制。")
    add_figure_placeholder(doc, "3-1", "漏洞响应闭环泳道图", "应使用横向泳道活动图。泳道包括普通用户、平台系统、审核员、管理员；节点包括完成准入、创建草稿、提交报告、形式校验、人工审核、要求补充或驳回、确认有效、修复跟踪、复测、归档、积分处理、证书签发与通知。各分支应标明状态变化及回退条件，箭头采用直线或直角折线。")

    add_heading(doc, "3.2.2 以可信服务边界提供按需AI辅助", 3)
    add_body(doc, "所有模型调用应由AI服务网关统一处理，客户端不得直接接触平台模型密钥。用户或审核员主动触发报告分析后，系统应提取必要文本，执行字段级去标识化、长度控制和权限校验，再向模型提供任务标准、最小必要上下文及结构化输出约束。AI结果应持久保存，并明确标识为报告完善或审核辅助建议。知识库独立保存经审核或规则准入的可复用内容，通过相似检索向后续问答提供有限上下文，以减少重复生成及无效模型令牌消耗；知识库条目不得创建或转换为Skill。")
    add_figure_placeholder(doc, "3-2", "AI辅助与知识检索泳道图", "应使用横向泳道活动图。泳道包括用户或审核员、客户端应用、AI服务网关、知识库、模型服务和审计服务；节点包括主动触发、权限与额度检查、文本解析、去标识化、知识检索、上下文组装、模型调用、结构化校验、结果持久化、用户反馈和审计。图中应明确表达“未主动触发时不调用模型”“AI不作最终业务决定”“外部服务故障不阻塞核心业务”。")

    add_heading(doc, "3.3 假设与依赖", 2)
    add_table(doc, "3-1", "假设与依赖", ["编号", "假设或依赖", "约束与应对"], [
        ["AD-01", "平台以统一认证码作为用户唯一身份标识；外部统一身份认证接入属于可配置集成能力。", "认证码应保持唯一；身份源切换或接入不得破坏既有用户及业务记录。"],
        ["AD-02", "平台依赖具备事务一致性的关系数据库保存业务数据，并依赖受控对象存储保存附件与证书文件。", "运行环境必须提供持久化存储、备份恢复、访问控制和容量监测能力。"],
        ["AD-03", "聊天即时消息依赖全双工实时通信通道及接入层升级转发能力。", "连接中断后应自动恢复并补拉持久化消息，不得以客户端内存作为唯一记录。"],
        ["AD-04", "邮件通知依赖管理员配置的可用邮件服务。", "发送失败不重试业务动作，但必须记录失败原因和审计事件。"],
        ["AD-05", "管理员与审核员只使用管理员配置的平台AI；普通用户可按平台策略使用平台模型或自有模型。", "客户端不得保存明文密钥；用户凭据应加密存储，同一服务提供方可配置多个凭据。"],
        ["AD-06", "外部模型的价格、质量、延迟与可用性不受平台保证。", "应设置超时、取消、额度和清晰错误反馈，核心漏洞流程不得依赖AI成功。"],
        ["AD-07", "部署方负责确定合法测试范围、授权资产、责任边界、漏洞等级和积分规则。", "系统提供配置与审计能力，不自动扩大授权范围。"],
        ["AD-08", "系统按单校园场景设计。", "系统范围不包含多学校租户隔离；相关扩展应另行开展需求分析与架构设计。"],
        ["AD-09", "生产部署具备HTTPS、反向代理、备份、持久存储和时间同步。", "缺少任一关键基础条件时不得按生产环境验收。"],
    ])


STAKEHOLDERS = [
    ["SH-01", "未登录访客", "查看公开介绍、规则、公告并验证证书真伪。", "无需登录即可完成必要公开查询，且不暴露敏感信息。"],
    ["SH-02", "普通用户", "完成安全准入，提交和跟踪本人漏洞，使用学习、积分、证书、通知、聊天和AI辅助。", "流程清晰、数据不丢失、关键状态及时感知、个人数据受保护。"],
    ["SH-03", "审核员", "审核漏洞、要求补充、提出结论、开展复测归档并管理必要公告或证书业务。", "队列清晰、材料完整、决策可追溯、AI只按需辅助。"],
    ["SH-04", "管理员", "维护人员权限、制度、资产、内容、证书、积分、AI、聊天与审计。", "统一配置、风险可控、异常可追踪、关键操作可撤销或有确认。"],
    ["SH-05", "系统运维人员", "部署、监控、备份、恢复和故障排查。", "配置明确、服务可观测、备份可恢复、外部依赖可诊断。"],
    ["SH-06", "资产责任单位", "根据组织流程接收漏洞处置信息并推动修复。", "仅获得履行处置职责所需的信息，不作为平台独立登录角色。"],
    ["SH-07", "外部服务提供方", "提供邮件、对象存储或模型服务。", "接口稳定、权限最小、故障隔离；不获得非必要业务数据。"],
]


USE_CASES = [
    {
        "id": "UC-01", "name": "公开访问与证书验证", "priority": "P0", "actors": "未登录访客、平台系统",
        "summary": "访客查看公开页面、规则和公告，并按证书编号验证证书真伪与状态。",
        "pre": "平台公开服务可访问；证书查询无需登录。",
        "trigger": "访客进入公开页面或提交符合格式的证书编号。",
        "main": ["系统展示未登录首页、平台规则、公告、关于平台与证书查询入口。", "访客输入证书编号并提交查询。", "系统校验证书编号格式、长度和请求频率。", "系统返回存在性、有效状态及公开核验所需的最少信息。", "访客可结束查询或继续查询其他编号。"],
        "alt": ["编号格式错误时在输入框附近提示，不发起无效请求。", "证书不存在、已撤销或服务异常时使用统一提示区分结果，不泄露持有人敏感信息。"],
        "post": "不修改证书业务数据；必要的查询安全日志可被记录。",
        "data": "证书编号、证书状态、签发日期、公开持有人脱敏信息、关联漏洞公开摘要。",
        "rules": ["BR-01-01 公开查询遵循最小披露原则。", "BR-01-02 被撤销证书必须明确显示无效状态。", "BR-01-03 输入必须满足证书编号格式、长度及频率限制。"],
    },
    {
        "id": "UC-02", "name": "注册、登录与密码找回", "priority": "P0", "actors": "普通用户、审核员、管理员、平台系统、邮件服务",
        "summary": "不同角色通过受控方式注册或登录，遗忘密码时通过已登记邮箱验证后重置。",
        "pre": "用户具备有效身份标识；邮件找回要求账号已绑定邮箱。",
        "trigger": "用户提交登录、注册或忘记密码请求。",
        "main": ["系统校验验证码、统一认证码或邮箱等输入。", "系统按当前注册策略执行开放注册或管理员审批。", "登录成功后签发访问令牌与刷新令牌并进入角色工作台。", "忘记密码时系统先确认邮箱存在关联账号，再发送一次性验证码。", "验证码有效且校验通过后，系统允许设置符合策略的新密码并撤销旧会话。"],
        "alt": ["账号不存在、状态异常、凭据错误或频率超限时返回统一且不过度暴露账号存在性的错误。", "邮件发送失败时不创建可被利用的半完成重置状态，并记录审计。"],
        "post": "认证会话、登录安全记录或密码重置记录被更新。",
        "data": "用户、角色、账号状态、密码摘要、邮箱、验证码、刷新令牌、登录IP和时间。",
        "rules": ["BR-02-01 密码不可逆散列存储。", "BR-02-02 验证码必须限时、限次、一次性使用。", "BR-02-03 管理员异地IP登录应产生安全提醒。"],
    },
    {
        "id": "UC-03", "name": "完成安全准入", "priority": "P0", "actors": "普通用户、管理员、平台系统",
        "summary": "普通用户阅读并签署安全责任书，通过网络安全知识答题后取得提交资格。",
        "pre": "用户已登录；管理员已配置有效责任书与启用题目。",
        "trigger": "用户首次登录收到准入提示，或主动开始准入流程。",
        "main": ["系统以弹窗展示责任书，用户阅读并确认签署。", "系统记录责任书版本、签署时间和用户身份。", "系统按题型规则生成试卷并展示单选、判断和多选题。", "用户提交答案，系统在服务端完成评分。", "达到默认80分或管理员配置阈值后，系统更新准入状态。"],
        "alt": ["用户尝试退出责任书或答题弹窗时，系统二次确认。", "题库不足、责任书失效或提交异常时不授予准入，并提供明确反馈。", "未通过时允许按平台规则重新答题。"],
        "post": "签署记录、考试记录、得分、题目版本和准入状态持久化。",
        "data": "责任书版本、签署记录、题库、试卷、答案、得分、通过阈值、准入状态。",
        "rules": ["BR-03-01 默认题型为10道单选、5道判断、5道多选，满分100分。", "BR-03-02 默认通过分数为80分，管理员可配置。", "BR-03-03 AI可辅助出题，但题目入库和启用由管理员决定。"],
    },
    {
        "id": "UC-04", "name": "提交漏洞报告", "priority": "P0", "actors": "普通用户、平台系统、对象存储、通知服务",
        "summary": "通过基础信息与Markdown正文提交单个漏洞，附件作为证据补充。",
        "pre": "用户已登录并完成安全准入；测试对象位于已授权范围。",
        "trigger": "用户新建报告或恢复未完成草稿。",
        "main": ["系统创建或恢复唯一未提交草稿。", "用户填写标题、影响资产、漏洞类型、责任单位等基础字段。", "系统把基础字段同步到Markdown报告，标题在文档中居中展示。", "用户编辑复现步骤、影响评估、修复建议和默认测试边界说明，可粘贴图片。", "用户上传原始报告或其他附件，系统执行类型、大小和安全校验。", "服务端执行必要字段与最低长度完整性校验。", "用户勾选知情与责任声明后提交。", "系统在事务中生成唯一漏洞编号、状态事件和通知。"],
        "alt": ["未通过完整性校验时定位到字段并保留草稿。", "附件上传失败时已填写正文不得丢失。", "粘贴图片失败时保留文本并使用统一提示。", "未授权资产或准入失效时拒绝提交并记录必要日志。"],
        "post": "形成正式漏洞报告、附件关联、初始状态、事件记录和提交通知。",
        "data": "草稿、漏洞编号、基础字段、Markdown正文、嵌入图片、附件、声明确认、提交时间。",
        "rules": ["BR-04-01 Markdown正文为主，附件为补充。", "BR-04-02 每次正式提交只对应一个漏洞。", "BR-04-03 影响评估和修复建议分别不少于10个字符。", "BR-04-04 原始附件不得因AI脱敏而被替换或修改。"],
    },
    {
        "id": "UC-05", "name": "恢复草稿与补充材料", "priority": "P0", "actors": "普通用户、平台系统、对象存储",
        "summary": "系统保存最后一个未提交草稿，并允许用户按审核要求补充文字和附件。",
        "pre": "用户拥有未提交草稿或本人漏洞处于允许补充的状态。",
        "trigger": "用户再次进入提交页，或在本人漏洞详情选择补充材料。",
        "main": ["再次进入提交页时，系统检测上次未完成草稿。", "系统弹窗询问恢复或放弃，不提供独立草稿管理模块。", "用户恢复后继续编辑，自动保存最新内容。", "收到补充要求时，用户进入本人漏洞详情。", "用户提交补充说明与附件，系统记录补充批次并通知审核方。"],
        "alt": ["用户选择放弃草稿时二次确认后删除。", "补充窗口已关闭或漏洞状态变化时拒绝写入并刷新详情。", "并发保存冲突时保留较新版本并提示用户。"],
        "post": "草稿被更新、删除或转为正式报告；补充材料形成不可混淆的时间记录。",
        "data": "草稿版本、保存时间、补充说明、补充附件、提交者和状态事件。",
        "rules": ["BR-05-01 仅保留用户最后一个未提交且未完成的草稿。", "BR-05-02 补充材料不覆盖原始报告与历史附件。", "BR-05-03 所有用户漏洞均保留补充材料入口，但服务端应依据业务状态决定是否接收。"],
    },
    {
        "id": "UC-06", "name": "查看漏洞进度与AI建议", "priority": "P1（AI建议为P2）", "actors": "普通用户、平台系统、AI服务网关、模型服务",
        "summary": "用户查看本人漏洞的紧凑信息、处理事件及按需生成的AI报告建议。",
        "pre": "用户已登录且报告属于本人。",
        "trigger": "用户进入“我的漏洞”或点击AI建议。",
        "main": ["系统列出漏洞编号、标题、状态、资产、类型、责任单位、时间、附件和积分。", "用户进入详情查看Markdown报告、附件和状态事件。", "用户点击AI建议时，系统跳转到AI助手并携带报告标识。", "AI服务网关提取并去标识化必要内容，调用与提交解析一致的评分标准。", "系统返回评分和修改建议，不出现再次提交卡片，并保存可追踪结果。"],
        "alt": ["用户未点击AI建议时不得发起模型调用。", "字段无法可靠脱敏时不向模型传递该字段，并在建议中说明限制。", "模型失败时保留原报告和页面状态。"],
        "post": "查看不改变漏洞状态；AI调用形成结果、消耗和审计记录。",
        "data": "本人漏洞摘要、事件、AI评分、建议、限制、模型运行记录。",
        "rules": ["BR-06-01 AI评分不代表漏洞有效性、风险等级或最终审核结论。", "BR-06-02 只有资源所有者及授权角色可查看报告。"],
    },
    {
        "id": "UC-07", "name": "审核漏洞并选择处置动作", "priority": "P0（AI辅助为P2）", "actors": "审核员、管理员、平台系统、AI服务网关",
        "summary": "审核员或管理员查看完整材料，按需获取AI审核建议，人工决定确认有效、要求补充或驳回。",
        "pre": "操作者具备审核权限；漏洞处于可审核状态。",
        "trigger": "操作者从审核队列进入漏洞详情。",
        "main": ["系统展示报告基础信息和Markdown正文。", "AI报告建议与附件材料以按钮打开二级弹窗，默认不展开。", "操作者可点击AI建议，系统将去标识化报告交由模型评分。", "AI以审核员辅助角色返回评分、依据、限制和建议动作，并将审核意见草稿填入对应输入框。", "操作者复核材料，可修改或删除AI生成的意见。", "操作者人工选择确认有效、要求补充或驳回，填写必要意见并确认。", "系统校验状态、权限和必填项，在事务中保存结论、事件、通知及审计。"],
        "alt": ["未点击AI建议时不消耗Token。", "已有持久化建议时直接展示；重新生成需主动触发并保留可审计版本。", "状态已被其他审核员改变时拒绝过期提交并要求刷新。"],
        "post": "审核结论与意见持久化；用户收到站内通知，重要状态按策略发送邮件。",
        "data": "完整报告、原始附件、AI建议、审核意见、审核动作、时间、操作者和审计日志。",
        "rules": ["BR-07-01 AI仅提出建议动作，最终动作必须由审核员或管理员确认。", "BR-07-02 AI意见填入后可编辑或删除，保存内容视为人工确认后的业务意见。", "BR-07-03 要求补充、驳回和确认有效必须通知提交者。"],
    },
    {
        "id": "UC-08", "name": "修复跟踪、复测与归档", "priority": "P0", "actors": "审核员、管理员、普通用户、平台系统",
        "summary": "对确认有效的漏洞跟踪修复，记录复测结果并完成归档。",
        "pre": "漏洞已确认有效并进入修复或复测阶段。",
        "trigger": "授权人员标记修复进展或提交复测结论。",
        "main": ["系统展示待修复或待复测队列。", "授权人员记录修复进展与必要说明。", "审核员执行复测并提交通过或未通过结论。", "未通过时系统回到修复流程并记录原因。", "通过时系统归档漏洞，固化最终时间线。", "系统通知相关用户，并触发后续积分或证书条件判断。"],
        "alt": ["材料不足或环境不可用时可记录暂缓原因，但不得伪造复测通过。", "并发状态冲突时拒绝提交过期结论。"],
        "post": "复测记录、最终状态、事件、通知和审计信息持久化。",
        "data": "修复说明、复测步骤、复测结果、附件、操作者、状态事件。",
        "rules": ["BR-08-01 复测记录不得覆盖审核记录。", "BR-08-02 归档前必须具备符合流程的有效结论。", "BR-08-03 归档属于重要通知事件。"],
    },
    {
        "id": "UC-09", "name": "签发、撤销与验证证书", "priority": "P0", "actors": "管理员、审核员、普通用户、未登录访客、平台系统",
        "summary": "授权人员依据已归档漏洞签发或撤销证书，用户下载，访客公开验证。",
        "pre": "关联漏洞满足证书签发规则；操作者具备证书权限。",
        "trigger": "授权人员在证书管理或漏洞详情中发起签发。",
        "main": ["系统校验关联漏洞、持有人和重复签发约束。", "操作者确认模板、签发信息与证书内容。", "系统生成唯一证书编号和受控证书文件。", "系统通知用户，用户可在“我的证书”查看与下载。", "访客可按UC-01验证证书。", "确需撤销时，授权人员填写原因并二次确认。"],
        "alt": ["模板或存储不可用时不创建半完成有效证书。", "重复签发或状态不符合时拒绝操作。", "撤销后下载和公开验证均显示最新状态。"],
        "post": "证书记录、文件、签发或撤销事件及通知持久化。",
        "data": "证书编号、持有人、关联漏洞、签发人、签发日期、状态、撤销原因、文件。",
        "rules": ["BR-09-01 证书编号全局唯一。", "BR-09-02 公开验证不得暴露不必要的个人信息和漏洞细节。", "BR-09-03 签发和撤销均属于高风险可审计操作。"],
    },
    {
        "id": "UC-10", "name": "积分、排行榜与商品兑换", "priority": "P1", "actors": "普通用户、管理员、平台系统",
        "summary": "平台维护积分账本、全校排行榜、商品与兑换处理。",
        "pre": "用户已登录；积分规则、商品库存和兑换限制已配置。",
        "trigger": "业务事件产生积分，或用户查看排行榜、发起兑换。",
        "main": ["系统按经过确认的业务规则写入积分账本。", "用户查看积分余额、明细及全校排行榜前100名。", "排行榜按北京时间计算，底部固定展示本人排名和积分。", "用户查看可兑换商品与限制，确认兑换。", "系统在事务中校验余额、库存、限购并扣减积分。", "管理员处理兑换完成或取消，系统同步状态与账本。"],
        "alt": ["积分不足、库存不足或超过限购时拒绝兑换。", "并发兑换冲突时只允许一笔成功。", "取消是否退分按商品和兑换规则执行并记录。"],
        "post": "积分账本、余额、排行榜数据和兑换状态一致。",
        "data": "积分账户、账本、排行榜快照或查询、商品、库存、限购、兑换记录。",
        "rules": ["BR-10-01 积分余额必须由可追溯账本支持，一致性目标100%。", "BR-10-02 用户不得退出公开排名。", "BR-10-03 排行榜仅提供全校榜，页面返回前100名。"],
    },
    {
        "id": "UC-11", "name": "学习中心与资源管理", "priority": "P1", "actors": "普通用户、管理员、平台系统、对象存储",
        "summary": "用户浏览和下载学习资源，管理员维护资源内容、封面与可用状态。",
        "pre": "用户已登录；管理员具备学习资源管理权限。",
        "trigger": "用户进入学习中心或管理员进入学习管理。",
        "main": ["系统按状态和分类展示可用资源。", "用户查看资源并发起下载或访问。", "系统记录有效下载次数用于用户活动分析。", "管理员创建、编辑、启用、停用或删除资源。", "管理员上传封面或文件时，系统执行安全校验并受控存储。"],
        "alt": ["资源停用或文件不存在时阻止访问并提示。", "上传失败不保存指向无效对象的正式资源。"],
        "post": "资源及下载统计更新；管理操作进入审计日志。",
        "data": "资源标题、类别、描述、文件或链接、封面、状态、下载记录。",
        "rules": ["BR-11-01 学习活动指标中的“学习”按用户在学习页下载资源次数采集。", "BR-11-02 删除受引用资源前应进行影响确认。"],
    },
    {
        "id": "UC-12", "name": "公告、通知与受控聊天", "priority": "P0/P1", "actors": "普通用户、审核员、管理员、平台系统、邮件服务",
        "summary": "提供公告、业务通知、未读提示、已读状态和角色间受控即时聊天。",
        "pre": "用户已登录；聊天需具备chat:use权限；邮件服务按策略配置。",
        "trigger": "业务事件产生通知，管理员发布公告，或用户发起会话。",
        "main": ["系统在通知中心统一展示平台公告、系统通知、私聊和群聊。", "重要业务通知同步形成铃铛及相关菜单红点。", "用户阅读后，系统更新通知、实体和频道已读状态。", "用户可直接联系管理员和审核员，并按统一认证码搜索其他用户。", "普通用户首次发起会话，在对方回复前最多发送3条文字消息。", "管理员可创建普通群、全员群并管理成员；其他角色仅发言。", "WebSocket即时推送消息，断线后通过持久化记录补拉。"],
        "alt": ["敏感词命中时阻止发送、告知原因并记录审计。", "无chat:use权限或会话被关闭时禁止发送。", "邮件失败不回滚业务事件，不自动重试，但必须写入审计。"],
        "post": "公告、通知、已读状态、会话、消息和审计记录持久化。",
        "data": "公告、通知渠道、实体关联、会话、成员、文字消息、已读游标、敏感词、发送失败记录。",
        "rules": ["BR-12-01 普通群人数上限200人，全员群不受此普通群上限约束。", "BR-12-02 只有管理员可建群、加人、踢人、改群名和解散。", "BR-12-03 聊天仅支持文字；重要状态邮件只包含编号、标题和脱敏摘要。", "BR-12-04 聊天存储达到数据库容量80%时提醒管理员清理并支持导出备份。"],
    },
    {
        "id": "UC-13", "name": "使用AI助手、模型与声明式Skill", "priority": "P2", "actors": "普通用户、管理员、审核员、AI服务网关、模型服务",
        "summary": "用户通过会话使用AI助手、模型配置与声明式Skill；管理员维护平台模型和系统Skill。",
        "pre": "用户已登录；存在符合角色策略的可用模型；AI服务可访问。",
        "trigger": "用户主动发送问题、上传报告解析或选择已固定Skill。",
        "main": ["用户在模型配置弹窗选择平台模型或配置自有服务凭据。", "系统加密保存凭据并提供连通性测试，不回显明文密钥。", "用户在Skill管理弹窗查看公开Skill、固定最多3个并调整顺序。", "用户发送消息时，系统记录本轮明确启用的Skill。", "AI服务网关组织最小必要上下文、检索可用知识并调用所选模型。", "系统在同一历史会话中允许切换模型后继续对话。", "用户可对回复点赞或点踩，点踩内容进入人工审核条件。"],
        "alt": ["模型鉴权失败、超时或取消时返回可理解原因，不丢失用户消息。", "模型能力不足由用户自行承担选择结果，平台不做能力等级检测。", "管理员和审核员只能使用管理员配置的平台AI。"],
        "post": "会话、消息、模型选择、Skill使用、运行消耗和反馈持久化。",
        "data": "会话、消息、模型供应商、加密凭据、模型配置、Skill内容与固定关系、运行记录、反馈。",
        "rules": ["BR-13-01 同一轮最多选择3个Skill。", "BR-13-02 用户Skill为声明式内容，公开申请必须人工审核，AI不审核。", "BR-13-03 Skill不授予额外RBAC权限。", "BR-13-04 无点赞不阻止自动知识准入，但点踩必须转人工审核。"],
    },
    {
        "id": "UC-14", "name": "用户、角色与权限管理", "priority": "P0", "actors": "管理员、平台系统",
        "summary": "管理员维护用户状态、角色、权限和注册审批，系统执行最小权限控制。",
        "pre": "管理员已登录并具备相应管理权限。",
        "trigger": "管理员查询用户、处理注册或变更权限。",
        "main": ["管理员按条件查询用户和待审批注册。", "管理员查看不包含秘密凭据的用户必要信息。", "管理员批准或拒绝注册申请。", "管理员按制度分配角色或权限点，并确认高风险变更。", "系统立即使新权限生效，必要时撤销旧会话。", "系统记录操作者、对象、前后值、时间和结果。"],
        "alt": ["禁止管理员移除维持系统治理所需的最后有效高权限主体。", "权限冲突、对象不存在或并发变更时拒绝操作。"],
        "post": "用户、角色、权限或注册状态更新，审计记录不可由普通业务角色修改。",
        "data": "用户、角色、权限点、角色权限、用户权限、注册申请、账号状态和审计记录。",
        "rules": ["BR-14-01 权限控制以服务端校验为准，隐藏客户端入口不能替代授权。", "BR-14-02 用户只可访问本人或明确授权的数据。", "BR-14-03 高权限变更应二次确认并100%审计。"],
    },
    {
        "id": "UC-15", "name": "管理AI知识库、平台模型与Skill审核", "priority": "P2", "actors": "管理员、平台系统、AI服务网关",
        "summary": "管理员在AI管理页审核知识、维护平台模型、管理系统Skill并人工审核用户Skill申请。",
        "pre": "管理员具备AI管理权限；相关候选知识或Skill申请已存在。",
        "trigger": "管理员进入AI管理页或收到未读管理提醒。",
        "main": ["系统按知识审核、系统Skill、用户Skill和运行配置分类展示数据。", "管理员查看已解决会话形成的知识候选、证据、相似项和反馈。", "管理员批准、驳回、编辑或维护知识条目及版本。", "经规则准入的高质量知识可自动入库并保留可审计记录；未通过规则的候选等待人工处理。", "管理员配置平台模型并测试连接。", "管理员创建或编辑系统Skill，控制是否向用户展示及是否默认固定。", "管理员人工审核用户Skill公开申请。"],
        "alt": ["点踩关联知识必须进入人工审核，不得自动发布。", "模型密钥未配置加密密钥时禁止保存并明确提示。", "自动入库失败时保留候选和审计，不创建Skill草案。"],
        "post": "知识、版本、审核结果、模型配置、Skill展示策略和审计记录更新。",
        "data": "知识条目、来源会话、相似度、反馈、版本、模型配置、加密凭据、系统Skill、用户Skill申请。",
        "rules": ["BR-15-01 知识库只用于相似问题检索与回答增强，不安装Skill、不创建Skill草案，也不能转换为Skill。", "BR-15-02 知识入库须满足可复用价值、质量和安全规则，并可被审计。", "BR-15-03 用户已解决提示仅针对至少2轮且有复用价值的会话。"],
    },
    {
        "id": "UC-16", "name": "系统设置、平台内容与审计", "priority": "P0（AI与外观配置为P2）", "actors": "管理员、平台系统、邮件服务、对象存储",
        "summary": "管理员统一维护平台外观、制度、题库、邮件、证书、配额与其他设置，并查询审计日志。",
        "pre": "管理员具备对应设置或审计权限。",
        "trigger": "管理员进入系统设置或审计日志页面。",
        "main": ["系统以左侧固定分类、右侧独立滚动区域展示设置。", "管理员修改平台名称、图标、字体、色卡和AI助手头像等外观。", "管理员编辑或上传Markdown安全责任书，维护题库和通过阈值。", "管理员配置邮件、证书模板、注册策略、AI额度与知识使用策略。", "上传文件时系统执行类型、大小和安全校验。", "管理员保存后系统使用统一全局提示反馈结果。", "管理员按操作者、类型、对象、IP和时间查询审计日志。"],
        "alt": ["配置校验失败时保留原生效值并定位错误。", "预览与正式应用分离；文件上传成功但配置保存失败时不得错误显示已生效。", "审计查询无权限时拒绝访问。"],
        "post": "有效配置按版本或更新时间持久化，关键修改进入审计。",
        "data": "平台设置、外观资源、字体、责任书、题库、邮件配置、证书模板、AI配额、审计日志。",
        "rules": ["BR-16-01 平台模型密钥和邮件凭据不得明文回显。", "BR-16-02 审计日志应脱敏，不记录密码、完整令牌、完整密钥或原始未脱敏漏洞正文。", "BR-16-03 重要配置变更应具备确认、结果反馈和可追溯性。"],
    },
]


RELATED_USE_CASES = {
    "UC-01": "UC-09 签发、撤销与验证证书",
    "UC-02": "UC-03 完成安全准入；UC-14 用户、角色与权限管理",
    "UC-03": "UC-02 注册、登录与密码找回；UC-04 提交漏洞报告；UC-16 系统设置、平台内容与审计",
    "UC-04": "UC-03 完成安全准入；UC-05 恢复草稿与补充材料；UC-07 审核漏洞并选择处置动作",
    "UC-05": "UC-04 提交漏洞报告；UC-07 审核漏洞并选择处置动作",
    "UC-06": "UC-04 提交漏洞报告；UC-13 使用AI助手、模型与声明式Skill",
    "UC-07": "UC-04 提交漏洞报告；UC-05 恢复草稿与补充材料；UC-08 修复跟踪、复测与归档",
    "UC-08": "UC-07 审核漏洞并选择处置动作；UC-09 签发、撤销与验证证书",
    "UC-09": "UC-01 公开访问与证书验证；UC-08 修复跟踪、复测与归档",
    "UC-10": "UC-08 修复跟踪、复测与归档；UC-09 签发、撤销与验证证书",
    "UC-11": "UC-02 注册、登录与密码找回",
    "UC-12": "UC-02 注册、登录与密码找回；UC-04至UC-09 漏洞响应与证书业务",
    "UC-13": "UC-06 查看漏洞进度与AI建议；UC-15 管理AI知识库、平台模型与Skill审核",
    "UC-14": "UC-02 注册、登录与密码找回；UC-16 系统设置、平台内容与审计",
    "UC-15": "UC-13 使用AI助手、模型与声明式Skill；UC-16 系统设置、平台内容与审计",
    "UC-16": "UC-03 完成安全准入；UC-09 签发、撤销与验证证书；UC-12 公告、通知与受控聊天",
}


def build_use_cases(doc: Document) -> None:
    add_heading(doc, "4 用例说明", 1)
    add_heading(doc, "4.1 系统参与者与总体用例", 2)
    add_body(doc, "系统以角色权限、资源归属和业务状态共同约束用例。资产责任单位属于业务协作对象，不设置独立登录角色；外部服务仅承担邮件、文件存储或模型调用等技术能力，不取得业务数据控制权和业务决策权。")
    add_heading(doc, "4.1.1 涉众汇总", 3)
    add_table(doc, "4-1", "涉众汇总", ["编号", "涉众", "主要职责或行为", "核心期望"], STAKEHOLDERS)
    add_heading(doc, "4.1.2 用例清单与优先级", 3)
    add_table(doc, "4-2", "用例清单与优先级", ["编号", "用例名称", "主要参与者", "优先级"], [[u["id"], u["name"], u["actors"], u["priority"]] for u in USE_CASES])
    add_figure_placeholder(doc, "4-1", "系统总体用例图", "应使用UML用例图。左侧参与者为未登录访客、普通用户；右侧参与者为审核员、管理员、系统运维人员及外部服务。系统边界内包含UC-01至UC-16，使用include关系表示身份认证、权限校验、通知、审计和附件存储等共性能力，使用extend关系表示AI按需辅助。不得把资产责任单位画成可直接登录角色。")

    for idx, use_case in enumerate(USE_CASES, 1):
        add_heading(doc, f"4.{idx + 1} {use_case['name']}", 2)
        add_use_case_table(doc, f"4-{idx + 2}", use_case, RELATED_USE_CASES[use_case["id"]])
        add_figure_placeholder(
            doc,
            f"4-{idx + 1}",
            f"{use_case['name']}活动图与系统顺序图",
            f"应分别绘制活动图和系统顺序图。活动图以“{use_case['trigger']}”为起点，覆盖基本流程、主要异常分支和最终状态；系统顺序图应包含参与者（{use_case['actors']}）、客户端应用、服务端业务服务、数据存储，以及本用例涉及的外部依赖。消息名称应与规约动作一致，权限拒绝、事务失败和通知处理应明确显示。",
        )


def build_nonfunctional(doc: Document) -> None:
    add_heading(doc, "5 非功能需求", 1)
    add_heading(doc, "5.1 性能需求", 2)
    add_table(doc, "5-1", "性能需求", ["编号", "需求", "验收方法"], [
        ["NFR-PERF-01", "系统应支持不少于300名用户同时保持登录状态。", "使用接近生产配置的环境进行并发会话和核心接口压测。"],
        ["NFR-PERF-02", "普通查询与列表接口在正常校园网络和目标负载下P95响应时间不超过2秒。", "对首页、列表、详情摘要和通知等代表接口采集服务端与端到端指标。"],
        ["NFR-PERF-03", "登录、漏洞提交和审核操作的P95响应时间不超过3秒，不含大文件实际传输时间。", "使用固定数据集与并发模型压测，排除外部邮件和AI等待。"],
        ["NFR-PERF-04", "常见校园网络下，未命中浏览器缓存的首屏达到可交互状态不超过4秒。", "在约定设备与网络节流条件下通过浏览器性能工具测量。"],
        ["NFR-PERF-05", "聊天消息在正常网络下从服务端接收到接收方可见的时间不超过2秒。", "使用两个独立会话端测量实时消息投递，并验证断线后的消息补拉。"],
        ["NFR-PERF-06", "AI任务不纳入普通接口时限；用户触发后2秒内必须看到处理中状态，并允许异步完成、取消或失败。", "模拟正常、超时、鉴权失败和服务不可用场景。"],
        ["NFR-PERF-07", "客户端应对非首屏页面及大型可视化、AI等资源实施按需加载，避免加载与当前任务无关的资源。", "分析发布构建产物和浏览器网络请求瀑布，并验证视觉质量及业务能力未降低。"],
    ])

    add_heading(doc, "5.2 可靠性需求", 2)
    add_table(doc, "5-2", "可靠性需求", ["编号", "需求", "验收方法"], [
        ["NFR-REL-01", "除计划维护外，月度服务可用性目标不低于99.5%。", "依据健康检查和外部监控记录计算。"],
        ["NFR-REL-02", "事务数据库以及漏洞报告、附件、证书和审计日志应至少每日备份。", "检查备份计划、备份产物、完整性校验和保留策略。"],
        ["NFR-REL-03", "系统RPO不超过24小时，RTO不超过4小时。", "每学期至少执行一次恢复演练并记录恢复点与耗时。"],
        ["NFR-REL-04", "提交、审核、积分扣减、证书签发等跨记录操作应保证事务一致性。", "通过故障注入验证不存在半完成业务状态。"],
        ["NFR-REL-05", "服务重启后，漏洞、通知、聊天、AI结果和其他持久记录不得依赖进程内存恢复。", "重启容器并核对前后数据与未读状态。"],
        ["NFR-REL-06", "AI、邮件或对象存储故障不得阻塞与其无关的核心业务。", "分别断开外部依赖并执行登录、查询和不依赖附件的核心操作。"],
        ["NFR-REL-07", "附件上传失败时，用户已填写的漏洞正文和基础字段不得丢失。", "中断上传、刷新或重新登录，验证草稿恢复。"],
    ])

    add_heading(doc, "5.3 易用性需求", 2)
    add_table(doc, "5-3", "易用性需求", ["编号", "需求", "验收方法"], [
        ["NFR-USA-01", "平台应适配桌面、平板和移动浏览器；普通用户核心流程可在移动端完成，复杂管理操作以桌面端为主。", "在代表性视口验证无内容遮挡、溢出和不可操作控件。"],
        ["NFR-USA-02", "Chrome、Edge、Firefox应支持当前及前两个主要版本。", "建立浏览器兼容性测试矩阵。"],
        ["NFR-USA-03", "明暗模式下文字、选项、弹窗、表格和状态提示均应清晰可辨。", "对核心页面执行视觉与对比度检查。"],
        ["NFR-USA-04", "加载、成功、警告和错误应统一使用平台全局反馈服务；字段错误应定位至对应输入。", "触发代表性成功与失败场景，检查样式、语义和可恢复动作。"],
        ["NFR-USA-05", "高风险操作和退出未完成责任书、答题或编辑流程时应二次确认。", "自动化和人工测试取消、确认及状态保留。"],
        ["NFR-USA-06", "核心交互应支持键盘操作、清晰焦点和可识别标签。", "使用键盘完成登录、查询、表单和弹窗主流程。"],
    ])

    add_heading(doc, "5.4 可维护性需求", 2)
    add_table(doc, "5-4", "可维护性需求", ["编号", "需求", "验收方法"], [
        ["NFR-MNT-01", "系统应按业务领域划分模块，身份授权、输入校验、反馈提示、文件处理和AI调用应通过统一边界复用。", "通过设计审查和源代码审查确认不存在绕过统一边界的重复实现。"],
        ["NFR-MNT-02", "数据库结构变更必须通过版本化迁移实施；发布过程应支持按变更范围更新客户端、服务端或完整系统，并支持独立重启与仅同步源代码。", "在受控验证环境演练各发布路径，检查执行结果、失败恢复和幂等性。"],
        ["NFR-MNT-03", "重要接口、状态机、权限和故障场景应具有自动化测试或明确回归清单。", "质量门禁执行类型检查、构建、单元或集成测试及关键端到端用例。"],
        ["NFR-MNT-04", "运行配置应通过受控配置机制管理，认证秘密、访问凭据和加密密钥不得写入版本化源代码。", "执行秘密信息扫描，并检查示例配置仅包含非敏感示例值。"],
        ["NFR-MNT-05", "日志应包含可关联请求编号，但不得记录密码、完整令牌、完整API Key或原始未脱敏漏洞正文。", "抽查成功、失败和异常日志。"],
    ])

    add_heading(doc, "5.5 可移植性需求", 2)
    add_table(doc, "5-5", "可移植性需求", ["编号", "需求", "验收方法"], [
        ["NFR-POR-01", "系统应支持在符合项目运行基线的64位Linux环境中采用标准容器编排方式部署。", "在满足基线条件的空白运行环境中，依据部署文档完成构建、数据迁移、启动和健康检查。"],
        ["NFR-POR-02", "系统接入层应支持HTTPS终止、HTTP请求转发、静态资源服务及实时通信协议升级，不得依赖特定反向代理产品。", "通过受控域名验证业务API、文件访问、客户端路由和实时消息连接。"],
        ["NFR-POR-03", "模型服务、邮件服务和对象存储的连接参数应可配置，不得与单一运行路径或服务提供方形成不可替换依赖。", "替换为符合接口契约的验证服务端点，并执行连接、异常和降级测试。"],
        ["NFR-POR-04", "客户端发布产物应能够由符合HTTP标准的Web服务托管。", "执行正式构建，并验证客户端路由回退、缓存控制和静态资源访问。"],
    ])

    add_heading(doc, "5.6 其他需求", 2)
    add_heading(doc, "5.6.1 安全与隐私", 3)
    add_table(doc, "5-6", "安全与隐私需求", ["编号", "需求", "验收方法"], [
        ["NFR-SEC-01", "系统应以RBAC、资源归属和状态前置条件实施服务端最小权限。", "执行跨角色、跨用户、直接调用API和ID枚举测试。"],
        ["NFR-SEC-02", "API应实施输入类型、格式、长度、文件、频率和恶意内容校验。", "使用异常边界值、恶意文件、XSS、路径与高频请求测试。"],
        ["NFR-SEC-03", "附件不得通过公开对象地址直接访问，下载和预览必须经过授权校验。", "尝试匿名访问、跨用户访问和过期链接访问。"],
        ["NFR-SEC-04", "平台API Key、邮件凭据等秘密应加密存储，并在界面和日志中掩码。", "检查数据库、接口响应、浏览器存储和日志。"],
        ["NFR-SEC-05", "发送外部AI前应进行字段级去标识化；无法可靠脱敏的参数不得发送。", "使用邮箱、手机号、身份标识、公网或内网资产等测试样本检查出站内容。"],
        ["NFR-SEC-06", "AI不得作出漏洞有效性、风险等级、审核动作和证书签发的最终决定。", "检查权限、接口与界面，确认关键动作始终需要授权人员确认。"],
        ["NFR-SEC-07", "重要操作应审计操作者、对象、动作、时间、IP、结果和必要的脱敏差异。", "抽查权限、审核、证书、配置、聊天拦截和邮件失败记录。"],
        ["NFR-SEC-08", "个人信息处理应遵循合法、正当、必要和最小范围原则。", "审查采集字段、公开查询、邮件内容、AI出站内容和保存期限。"],
    ])

    add_heading(doc, "5.6.2 数据保留与容量", 3)
    add_bullets(doc, [
        "漏洞报告、审核、复测、证书、积分和审计记录的保留期限应由部署单位依据适用制度配置；系统不得在缺少有效保留策略时自动删除相关记录。",
        "聊天数据容量达到数据库分配容量的80%时，系统应提醒管理员。管理员可按日期预览并清理历史记录，清理前可选择导出离线备份。",
        "删除和清理操作必须校验权限、显示影响范围、二次确认并记录审计；备份导出不是执行清理的强制前置条件。",
        "知识库应保留来源、版本、审核、反馈和维护记录，以支持撤回错误知识和分析检索效果。",
    ])

    add_heading(doc, "5.6.3 可观测性与时间", 3)
    add_bullets(doc, [
        "系统应提供应用服务、数据库、缓存、对象存储和AI依赖的健康检查，并区分核心服务与可降级服务。",
        "关键请求应具有请求编号；客户端与服务端错误反馈可向管理员提供该编号用于定位，但不得向普通用户暴露堆栈与内部路径。",
        "业务展示与排行榜统计采用北京时间；数据库时间宜统一保存为可明确转换的标准时间。",
        "监控应覆盖错误率、响应时间、任务积压、实时连接、存储容量、邮件失败和AI运行失败等指标。",
    ])


def build_appendix(doc: Document) -> None:
    add_heading(doc, "6 附件", 1)
    add_heading(doc, "6.1 需求优先级与范围说明", 2)
    add_table(doc, "6-1", "需求优先级与范围", ["优先级", "范围", "说明"], [
        ["P0", "认证、准入、漏洞、附件、审核、补充、复测、归档、证书、权限、关键通知、审计", "构成合法、安全、可追踪运行的核心闭环。"],
        ["P1", "积分、排行榜、商城、学习、公告、邮件、聊天", "增强参与反馈与业务协同，故障不应破坏P0数据。"],
        ["P2", "AI助手、报告解析、审核建议、知识库、模型配置、Skill、外观", "增强效率与体验，必须可降级且不取得业务最终决定权。"],
    ])

    add_heading(doc, "6.2 功能需求追踪矩阵", 2)
    trace_rows = []
    for u in USE_CASES:
        trace_rows.append([f"FR-{u['id'][3:]}", u["name"], u["id"], u["priority"], "功能测试、权限测试、异常流程测试"])
    add_table(
        doc,
        "6-2",
        "功能需求追踪矩阵",
        ["需求编号", "需求名称", "来源用例", "优先级", "主要验证方式"],
        trace_rows,
        widths=[1.4, 4.2, 1.8, 1.7, 4.8],
    )

    add_heading(doc, "6.3 图示索引", 2)
    figure_rows = [
        ["图3-1", "漏洞响应闭环泳道图", "总体业务流程", "展示漏洞从准入、提交、审核、修复、复测到归档、积分和证书的闭环。"],
        ["图3-2", "AI辅助与知识检索泳道图", "技术处理流程", "展示按需调用、去标识化、知识检索、模型处理、持久化、审计及降级。"],
        ["图4-1", "系统总体用例图", "系统范围与参与者", "展示UC-01至UC-16及主要参与者、外部服务和用例关系。"],
        ["图4-2", "公开访问与证书验证活动图、系统顺序图", "详细需求与测试设计", "覆盖公开访问、证书编号校验、查询结果及异常处理。"],
    ]
    for idx, u in enumerate(USE_CASES[1:], 3):
        figure_rows.append([f"图4-{idx}", f"{u['name']}活动图、系统顺序图", "详细需求与测试设计", "覆盖基本流程、异常分支、权限、事务、外部依赖和最终状态。"])
    add_table(doc, "6-3", "图示索引", ["图号", "图名", "主要用途", "主要内容"], figure_rows)

    add_heading(doc, "6.4 验收准则", 2)
    add_body(doc, "系统验收应覆盖需求追踪矩阵中的全部P0需求及合同或项目计划约定的P1、P2需求。验收内容包括角色权限、漏洞响应闭环、数据一致性、安全控制、故障降级、部署交付、备份恢复、兼容性和文档一致性。各项性能指标应在经批准的硬件配置、网络条件、数据规模和测试脚本下执行，并保存验证环境、执行时间、样本数据、原始结果和判定结论。")
    add_body(doc, "涉及角色边界、授权范围、数据保留、外部AI数据传输范围、核心状态机和安全控制的变更，应按照需求变更流程完成影响分析、审批、实现、回归测试和版本记录。未满足前置条件或缺少可验证证据的需求项不得判定为通过。")


def set_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def normalize_table_fonts(doc: Document) -> None:
    for table in doc.tables:
        for run in table._tbl.iter(qn("w:r")):
            run_properties = run.find(qn("w:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("w:rPr")
                run.insert(0, run_properties)
            fonts = run_properties.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                run_properties.insert(0, fonts)
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), "宋体")
            for tag in ("sz", "szCs"):
                size = run_properties.find(qn(f"w:{tag}"))
                if size is None:
                    size = OxmlElement(f"w:{tag}")
                    run_properties.append(size)
                size.set(qn("w:val"), "21")


def normalize_table_paragraphs(doc: Document) -> None:
    """Remove inherited body indentation and list numbering inside tables."""
    for table_index, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        is_use_case = bool(table.rows) and table.rows[0].cells[0].text.strip() == "用例名"
        for row in table.rows:
            allow_automatic_row_height(row)
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, top=72, start=90, bottom=72, end=90)
                set_cell_shading(cell, "FFFFFF")

        if table_index == 0:
            for row in table.rows:
                set_cell_shading(row.cells[0], "E7E7E7")
        elif is_use_case:
            for row in table.rows[:6]:
                set_cell_shading(row.cells[0], "E7E7E7")
        elif table.rows:
            for cell in table.rows[0].cells:
                set_cell_shading(cell, "D9D9D9")

        for paragraph in table._tbl.iter(qn("w:p")):
            paragraph_properties = paragraph.find(qn("w:pPr"))
            if paragraph_properties is None:
                paragraph_properties = OxmlElement("w:pPr")
                paragraph.insert(0, paragraph_properties)

            paragraph_style = paragraph_properties.find(qn("w:pStyle"))
            if paragraph_style is None:
                paragraph_style = OxmlElement("w:pStyle")
                paragraph_properties.insert(0, paragraph_style)
            # Reference-table XML may point at styles that do not exist in the
            # generated package. LibreOffice can detach those paragraphs from
            # a table when repaginating a multi-page table.
            paragraph_style.set(qn("w:val"), "Normal")

            numbering = paragraph_properties.find(qn("w:numPr"))
            if numbering is not None:
                paragraph_properties.remove(numbering)

            indentation = paragraph_properties.find(qn("w:ind"))
            if indentation is None:
                indentation = OxmlElement("w:ind")
                paragraph_properties.append(indentation)
            # Explicit zero values override the two-character indentation
            # inherited from the document's Normal style in Word/WPS.
            for attribute in ("left", "right", "start", "end", "firstLine", "firstLineChars"):
                indentation.set(qn(f"w:{attribute}"), "0")
            for attribute in ("hanging", "hangingChars"):
                indentation.attrib.pop(qn(f"w:{attribute}"), None)

            spacing = paragraph_properties.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                paragraph_properties.append(spacing)
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "240")
            spacing.set(qn("w:lineRule"), "auto")


def deduplicate_package(path: Path) -> None:
    """Remove duplicate ZIP members inherited from the legacy .doc conversion."""
    with ZipFile(path, "r") as source:
        names = list(dict.fromkeys(source.namelist()))
        payloads = {name: source.read(name) for name in names}
    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        with ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
            for name in names:
                target.writestr(name, payloads[name])
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)


def main() -> None:
    if not TEMPLATE.exists():
        raise SystemExit(f"Template not found: {TEMPLATE}")
    doc = Document(TEMPLATE)
    remove_body_content(doc)
    configure_styles(doc)
    configure_section(doc)
    set_update_fields_on_open(doc)

    build_cover(doc)
    build_document_control(doc)
    build_toc(doc)
    build_intro(doc)
    build_business(doc)
    build_vision(doc)
    build_use_cases(doc)
    build_nonfunctional(doc)
    build_appendix(doc)
    normalize_table_fonts(doc)
    normalize_table_paragraphs(doc)

    doc.core_properties.title = f"{PROJECT_NAME}软件需求规格说明书"
    doc.core_properties.subject = "软件需求规格说明书"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "校园漏洞响应, 安全学习, 软件需求规格说明书"
    doc.core_properties.comments = "正式软件需求规格说明书。"
    doc.save(OUTPUT)
    deduplicate_package(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
